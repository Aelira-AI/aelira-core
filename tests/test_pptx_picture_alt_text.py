"""PPTX picture alt text lives on the picture's non-visual properties.

In PresentationML the `descr` attribute belongs to `p:nvPicPr/p:cNvPr`, not to
the `p:pic` element. The scanner must read it from there, the same place the
remediator writes it, so pictures that already have alt text are not reported
and a saved-file rescan sees the alt text remediation wrote.
"""

import io
import zipfile

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation

from src.education.pptx_processor import PowerPointProcessor, get_shape_alt_text
from src.education.remediation.base import IssueCategory, RemediationConfig
from src.education.remediation.pptx_remediator import PptxRemediator

pytestmark = pytest.mark.unit

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NS = {"p": P_NS}

DESCRIBED = "Bar chart of enrolment by year, rising from 120 to 480 students"
GENERATED = "Line graph of weekly rainfall peaking in March"


def _png(colour: str = "green") -> io.BytesIO:
    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), colour).save(buffer, format="PNG")
    buffer.seek(0)
    return buffer


def _set_descr(picture, descr):
    """Put alt text where PowerPoint does: <p:cNvPr ... descr="..."/>.

    python-pptx fills `descr` with the source file name on every picture it
    adds, so `None` has to remove that default to leave a picture undescribed.
    """
    properties = picture._element.nvPicPr.cNvPr
    if descr is None:
        properties.attrib.pop("descr", None)
    else:
        properties.set("descr", descr)


def _build(path, pictures):
    """Save a one-slide deck. `pictures` is a list of (name, descr or None)."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Enrolment"
    for index, (name, descr) in enumerate(pictures):
        picture = slide.shapes.add_picture(_png(), index * 914400, 1828800)
        picture.name = name
        _set_descr(picture, descr)
    presentation.save(path)


def _alt_findings(result):
    return [issue for issue in result.issues if issue["category"] == "alt_text"]


def _slide_xml(path):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read("ppt/slides/slide1.xml"))


def _picture_descr(path, name):
    """The `descr` on the named picture's cNvPr in the saved package."""
    props = _slide_xml(path).xpath(
        "//p:pic/p:nvPicPr/p:cNvPr[@name=$name]", namespaces=NS, name=name
    )
    assert len(props) == 1, f"expected one picture named {name!r}"
    return props[0].get("descr")


class _VisionStub:
    """Purpose-bound alt-text client that always returns a description."""

    def __init__(self, text=GENERATED):
        self.text = text
        self.calls = 0

    def analyze_image_sync(self, *args, **kwargs):
        self.calls += 1
        return {"success": True, "content": self.text}


def _config(tmp_path):
    return RemediationConfig(
        use_ai=True,
        allow_legacy_nested_ai=False,
        create_backup=False,
        output_directory=str(tmp_path / "output"),
    )


class TestScannerReadsCNvPr:
    def test_picture_with_alt_text_is_not_reported(self, tmp_path):
        source = tmp_path / "described.pptx"
        _build(source, [("Enrolment chart", DESCRIBED)])

        result = PowerPointProcessor().process_pptx(str(source))

        assert result.total_images == 1
        assert [i for s in result.slides for i in s.alt_text_issues] == []
        assert _alt_findings(result) == []

    def test_picture_without_alt_text_is_still_reported(self, tmp_path):
        source = tmp_path / "undescribed.pptx"
        _build(source, [("Mystery photo", None)])

        result = PowerPointProcessor().process_pptx(str(source))

        findings = [i for s in result.slides for i in s.alt_text_issues]
        assert len(findings) == 1
        assert findings[0].shape_name == "Mystery photo"
        assert findings[0].has_alt_text is False
        assert len(_alt_findings(result)) == 1

    def test_whitespace_only_alt_text_is_reported_as_missing(self, tmp_path):
        source = tmp_path / "blank.pptx"
        _build(source, [("Blank descr", "   ")])

        result = PowerPointProcessor().process_pptx(str(source))

        assert [i.has_alt_text for s in result.slides for i in s.alt_text_issues] == [
            False
        ]

    def test_only_the_undescribed_picture_in_a_mixed_slide_is_reported(self, tmp_path):
        source = tmp_path / "mixed.pptx"
        _build(source, [("Has alt", DESCRIBED), ("Needs alt", None)])

        result = PowerPointProcessor().process_pptx(str(source))

        findings = [i for s in result.slides for i in s.alt_text_issues]
        assert [f.shape_name for f in findings] == ["Needs alt"]
        assert [f["description"] for f in _alt_findings(result)] == [
            "Image 'Needs alt' is missing alt text"
        ]

    def test_fully_described_deck_is_not_capped_by_a_critical_finding(self, tmp_path):
        """Missing alt text is critical and caps the score at 49."""
        described = tmp_path / "described.pptx"
        undescribed = tmp_path / "undescribed.pptx"
        _build(described, [("Chart", DESCRIBED)])
        _build(undescribed, [("Chart", None)])

        with_alt = PowerPointProcessor().process_pptx(str(described))
        without_alt = PowerPointProcessor().process_pptx(str(undescribed))

        assert not any(i["severity"] == "critical" for i in with_alt.issues)
        assert any(i["severity"] == "critical" for i in without_alt.issues)
        assert without_alt.compliance_score <= 49
        assert with_alt.compliance_score > 49

    def test_alt_text_on_the_pic_element_itself_does_not_count(self, tmp_path):
        """`descr` on p:pic is not where PresentationML keeps alt text."""
        source = tmp_path / "wrong_element.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = "Enrolment"
        picture = slide.shapes.add_picture(_png(), 0, 1828800)
        _set_descr(picture, None)
        picture._element.set("descr", DESCRIBED)
        presentation.save(source)

        result = PowerPointProcessor().process_pptx(str(source))

        assert [i.has_alt_text for s in result.slides for i in s.alt_text_issues] == [
            False
        ]


class TestGetShapeAltText:
    def _slide(self):
        presentation = Presentation()
        return presentation.slides.add_slide(presentation.slide_layouts[6])

    def test_picture(self):
        slide = self._slide()
        picture = slide.shapes.add_picture(_png(), 0, 0)
        _set_descr(picture, None)
        assert get_shape_alt_text(picture) == ""
        _set_descr(picture, DESCRIBED)
        assert get_shape_alt_text(picture) == DESCRIBED

    def test_autoshape_and_textbox(self):
        slide = self._slide()
        box = slide.shapes.add_textbox(0, 0, 914400, 914400)
        box._element.nvSpPr.cNvPr.set("descr", "A caption box")
        assert get_shape_alt_text(box) == "A caption box"

    def test_group_reads_its_own_properties_not_a_childs(self):
        slide = self._slide()
        group = slide.shapes.add_group_shape()
        child = group.shapes.add_picture(_png(), 0, 0)
        _set_descr(child, "Child picture")
        assert get_shape_alt_text(group) == ""
        group._element.nvGrpSpPr.cNvPr.set("descr", "The whole group")
        assert get_shape_alt_text(group) == "The whole group"
        assert get_shape_alt_text(child) == "Child picture"

    def test_graphic_frame(self):
        slide = self._slide()
        table = slide.shapes.add_table(1, 1, 0, 0, 914400, 914400)
        table._element.nvGraphicFramePr.cNvPr.set("descr", "A one-cell table")
        assert get_shape_alt_text(table) == "A one-cell table"


class TestRemediateThenRescan:
    def _remediate(self, source, tmp_path, client):
        scan = PowerPointProcessor().process_pptx(str(source))
        issues = _alt_findings(scan)
        assert issues, "scanner reported no missing alt text for the fixture"
        result = PptxRemediator(
            str(source), issues, _config(tmp_path), alt_text_client=client
        ).remediate()
        assert result.success, result.error_message
        return scan, result

    def test_written_alt_text_passes_the_saved_file_rescan(self, tmp_path):
        source = tmp_path / "source.pptx"
        _build(source, [("Rainfall chart", None)])
        client = _VisionStub()

        before, result = self._remediate(source, tmp_path, client)

        # The remediator wrote the description onto the picture's cNvPr.
        assert client.calls == 1
        assert _picture_descr(result.output_file, "Rainfall chart") == GENERATED
        assert (
            _slide_xml(result.output_file).xpath("//p:pic[@descr]", namespaces=NS) == []
        )

        # The saved-file rescan sees it, so the fix is verified, not demoted.
        alt_fixed = [
            fix for fix in result.fixed_issues if fix.category == IssueCategory.ALT_TEXT
        ]
        assert len(alt_fixed) == 1
        assert alt_fixed[0].verification_passed is True
        assert result.manual_issues == []
        assert result.verification_passed is True
        assert result.verification_result.regressions == []

        after = PowerPointProcessor().process_pptx(result.output_file)
        assert _alt_findings(after) == []
        assert after.compliance_score > before.compliance_score
        assert result.remediated_compliance_score == after.compliance_score

    def test_existing_alt_text_is_left_alone(self, tmp_path):
        source = tmp_path / "source.pptx"
        _build(source, [("Has alt", DESCRIBED), ("Needs alt", None)])
        client = _VisionStub()

        _, result = self._remediate(source, tmp_path, client)

        assert client.calls == 1
        assert _picture_descr(result.output_file, "Has alt") == DESCRIBED
        assert _picture_descr(result.output_file, "Needs alt") == GENERATED
        assert result.verification_passed is True
        assert result.manual_issues == []
