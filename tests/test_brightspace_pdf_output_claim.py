"""Brightspace PDF promotion preserves descriptor-bound output ownership."""

from __future__ import annotations

import asyncio
import builtins
import hashlib
import os
import threading
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from src.db.models import CloudProvider, Scan, ScanFix, ScanStatus
from src.education.remediation.base import RemediationResult
from src.education.remediation.output_claim import DescriptorBoundOutputClaim
from src.services.remediation_artifact_service import ArtifactPublicationResult
from src.services.scan_fix_service import ScanReviewGraph
from tests.test_image_equation_review_gate import EVIDENCE, _fix as _typed_equation_fix

CLAIMED_PDF = b"%PDF-1.7\nbrightspace descriptor claim\n%%EOF\n"


class _PdfResult(SimpleNamespace):
    def __init__(
        self,
        output_path: Path,
        *,
        with_claim: bool = True,
        metadata_error: Exception | None = None,
        metadata_override: dict[str, object] | None = None,
        fail_stream_on: int | None = None,
    ) -> None:
        super().__init__(
            success=True,
            output_file=str(output_path),
            verification_passed=True,
            fixed_count=1,
            manual_count=0,
            failed_count=0,
        )
        self.claim = None
        self.close_calls = 0
        self.stream_calls = 0
        self.metadata_error = metadata_error
        self.metadata_override = metadata_override
        self.fail_stream_on = fail_stream_on
        if with_claim:
            descriptor = os.open(output_path, os.O_RDONLY)
            self.claim = DescriptorBoundOutputClaim._snapshot_from_owned_descriptor(
                descriptor,
                filename=output_path.name,
                display_path=str(output_path),
                mime="application/pdf",
            )

    def has_output_claim(self) -> bool:
        return self.claim is not None and not self.claim.closed

    def output_claim_metadata(self) -> dict[str, object]:
        if self.metadata_error is not None:
            raise self.metadata_error
        if not self.has_output_claim():
            raise RuntimeError("missing output claim")
        if self.metadata_override is not None:
            return self.metadata_override
        return {
            "size_bytes": self.claim.size,
            "sha256": self.claim.sha256,
            "mime_type": self.claim.mime,
            "filename": self.claim.filename,
        }

    @contextmanager
    def open_output_stream(self):
        self.stream_calls += 1
        if self.fail_stream_on == self.stream_calls:
            if self.claim is not None:
                self.claim.close()
            raise RuntimeError("output claim closed during stream handoff")
        if not self.has_output_claim():
            raise RuntimeError("missing output claim")
        with self.claim.open_stream() as stream:
            yield stream

    def close_output_claim(self) -> None:
        self.close_calls += 1
        if self.claim is not None:
            self.claim.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("manual,failed", [(1, 0), (0, 1), (2, 1)])
async def test_verified_partial_pdf_remains_downloadable_with_unresolved_counts(
    tmp_path, manual, failed
):
    output = tmp_path / "partial.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    result.manual_count = manual
    result.failed_count = failed
    run = await _run_pdf_outer(tmp_path, result)
    assert run.outcome.status == "manual_required"
    assert run.outcome.has_remediated_version is True
    assert run.outcome.artifact_id == "artifact-pdf"
    assert run.outcome.fixed_count == 1
    assert run.outcome.manual_count == manual
    assert run.outcome.failed_count == failed
    assert run.cloud_file.remediated_issues_remaining == manual + failed
    assert run.cloud_file.writeback_status == "pending_review"
    assert run.published == [CLAIMED_PDF]
    assert run.validated == [CLAIMED_PDF]
    assert result.close_calls == 1


@pytest.mark.asyncio
async def test_partial_pdf_does_not_bypass_failed_verification(tmp_path):
    output = tmp_path / "partial.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    result.manual_count = 1
    result.verification_passed = False
    run = await _run_pdf_outer(tmp_path, result)
    assert run.outcome.has_remediated_version is False
    assert run.published == []
    assert result.close_calls == 1


@pytest.mark.asyncio
async def test_partial_pdf_uses_authenticated_download_route(tmp_path):
    from io import BytesIO
    from src.api.education import remediation_routes as routes

    output = tmp_path / "partial.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    result.manual_count = 2
    run = await _run_pdf_outer(tmp_path, result)
    artifact = _artifact()
    artifact.filename = "partial.pdf"
    scan = SimpleNamespace(id="scan-pdf", remediation_outcome="manual_required")
    principal = SimpleNamespace(department_id="dept-1")
    service = MagicMock()

    @contextmanager
    def verified(_db, candidate, **authority):
        assert candidate is artifact
        assert authority["department_id"] == "dept-1"
        assert authority["scan_id"] == "scan-pdf"
        assert authority["cloud_file_id"] == run.cloud_file.id
        with BytesIO(run.published[0]) as stream:
            yield stream

    service.open_verified.side_effect = verified
    run.db.query.return_value.filter.return_value.first.return_value = None
    with (
        patch.object(
            routes,
            "_managed_artifact_authority",
            return_value=(scan, run.cloud_file, artifact),
        ) as authorize,
        patch.object(
            routes.RemediationArtifactService, "from_settings", return_value=service
        ),
    ):
        response = await routes.download_managed_artifact(
            "scan-pdf", artifact.id, run.db, principal
        )
        downloaded = b"".join([chunk async for chunk in response.body_iterator])
    assert downloaded == CLAIMED_PDF
    authorize.assert_called_once_with(
        run.db, scan_id="scan-pdf", artifact_id=artifact.id, principal=principal
    )
    assert response.media_type == "application/pdf"
    assert run.cloud_file.remediated_issues_remaining == 2


def test_partial_artifact_job_keeps_manual_state_without_failed_job():
    from src.api.brightspace_routes import RemediationOutcome
    from src.jobs.brightspace_content_job import (
        _public_outcome,
        _commit_terminal_outcome,
    )

    outcome = RemediationOutcome(
        cloud_file_id="cloud-pdf",
        status="manual_required",
        fixed_count=1,
        manual_count=2,
        failed_count=1,
        has_remediated_version=True,
        artifact_id="artifact-pdf",
    )
    public = _public_outcome(outcome, scan_id="scan-pdf")
    assert public["status"] == "manual_required"
    assert public["download_available"] is True
    db = MagicMock()
    job = SimpleNamespace(id="job", claim_token="claim", worker_id="worker")
    owner = SimpleNamespace()
    db.execute.return_value.scalar_one_or_none.return_value = owner
    _commit_terminal_outcome(db, job, public)
    assert owner.status == "completed"
    assert owner.result_data["manual_count"] == 2
    assert owner.result_data["failed_count"] == 1
    assert "manual review" in owner.progress_message


def _cloud_file() -> SimpleNamespace:
    return SimpleNamespace(
        id="cloud-pdf",
        department_id="dept-1",
        last_scan_id="scan-pdf",
        provider=CloudProvider.BRIGHTSPACE.value,
        provider_metadata={
            "org_unit_id": 42,
            "url": "/content/source.pdf",
            "topic_type": "file",
        },
        file_name="source.pdf",
        file_size_bytes=len(CLAIMED_PDF),
        provider_file_id="7",
        content_body=None,
        remediated_body=None,
        has_remediated_version=False,
        remediation_origin=None,
        remediated_issues_fixed=0,
        remediated_issues_remaining=0,
        writeback_status=None,
    )


def _db(*, commit_error: Exception | None = None) -> MagicMock:
    scan_result = SimpleNamespace(
        issues=[
            {
                "id": "pdf-tagged",
                "category": "structure",
                "severity": "high",
                "description": "PDF is untagged",
            }
        ]
    )
    db = MagicMock()
    query = MagicMock()
    query.filter.return_value = query
    query.first.return_value = scan_result
    scan = SimpleNamespace(
        id="scan-pdf",
        department_id="dept-1",
        document_source="cloud_file",
        document_id="cloud-pdf",
        file_hash=hashlib.sha256(b"%PDF-source").hexdigest(),
        current_remediation_artifact_id=None,
        scan_type="PDF",
        status=ScanStatus.COMPLETED,
        remediation_outcome=None,
    )
    scan_query = MagicMock()
    scan_query.filter.return_value = scan_query
    scan_query.with_for_update.return_value = scan_query
    scan_query.populate_existing.return_value = scan_query
    scan_query.first.return_value = scan

    def query_for(model):
        if model is Scan:
            return scan_query
        if model is ScanFix.id:
            return db.query.return_value
        return query

    db.query.side_effect = query_for
    db.source_scan = scan
    if commit_error is not None:
        db.commit.side_effect = commit_error
    return db


def _artifact() -> SimpleNamespace:
    return SimpleNamespace(
        id="artifact-pdf",
        scan_id="scan-pdf",
        department_id="dept-1",
        cloud_file_id="cloud-pdf",
        provider_result={},
        lifecycle_status="available",
        mime_type="application/pdf",
        size_bytes=len(CLAIMED_PDF),
        sha256=hashlib.sha256(CLAIMED_PDF).hexdigest(),
        expires_at=datetime.now(timezone.utc) + timedelta(hours=1),
        review_status="pending",
        written_back_at=None,
        approved_at=None,
        approved_by_id=None,
        approved_by_ref=None,
        approval_checksum=None,
        approval_review_digest=None,
        rejected_by_id=None,
        rejected_by_ref=None,
        rejected_at=None,
    )


async def _run_pdf_outer(
    tmp_path: Path,
    result: _PdfResult,
    *,
    tamper: str | None = None,
    publication_error: BaseException | None = None,
    publication_hook=None,
    commit_error: Exception | None = None,
    matterhorn_result=None,
    matterhorn_error: Exception | None = None,
    source_bytes: bytes = b"%PDF-source",
    source_hash: str | None = "default",
    final_source_hash: str | None = "default",
):
    from src.api.brightspace_routes import (
        _WorkerRemediationResult,
        _remediate_file_impl,
    )

    cloud_file = _cloud_file()
    db = _db(commit_error=commit_error)
    db.source_scan.file_hash = (
        hashlib.sha256(source_bytes).hexdigest()
        if source_hash == "default"
        else source_hash
    )
    api_client = AsyncMock()
    api_client.get_topic_file.return_value = (source_bytes, "application/pdf")
    service = MagicMock()
    published: list[bytes] = []
    artifacts = []

    def publish(_db_arg, **kwargs):
        published.append(kwargs["source_stream"].read())
        if publication_hook is not None:
            publication_hook()
        if publication_error is not None:
            raise publication_error
        artifact = _artifact()
        artifact.sha256 = kwargs["claimed_sha256"]
        artifact.size_bytes = kwargs["claimed_size_bytes"]
        artifact.provider_result = kwargs["provider_result"]
        cloud_file.current_remediation_artifact_id = artifact.id
        artifacts.append(artifact)
        if final_source_hash != "default":
            db.source_scan.file_hash = final_source_hash
        return artifact

    service.claim_and_publish_stream.side_effect = publish
    service.claim_and_publish.side_effect = AssertionError(
        "Brightspace PDF publication must not use a pathname"
    )
    validated: list[bytes] = []
    validator = MagicMock()

    def validate(path):
        validation_path = Path(path)
        assert validation_path != Path(result.output_file)
        assert validation_path.is_file()
        validated.append(validation_path.read_bytes())
        if matterhorn_error is not None:
            raise matterhorn_error
        if matterhorn_result is not None:
            return matterhorn_result
        checkpoint = SimpleNamespace(
            id="01-003",
            name="Structure tree present",
            status=SimpleNamespace(value="pass"),
            severity="error",
            details=None,
            page_number=None,
        )
        return SimpleNamespace(
            checkpoints=[checkpoint], passed=1, failed=0, warnings=0, total=1
        )

    validator.validate.side_effect = validate

    async def run_worker(_department_id, worker, *args, **kwargs):
        assert worker.__name__ == "_run_remediator_worker"
        output = Path(result.output_file)
        if tamper == "replace":
            replacement = output.with_suffix(".replacement")
            replacement.write_bytes(b"%PDF-replaced-path")
            replacement.replace(output)
        elif tamper == "truncate":
            output.write_bytes(b"")
        elif tamper == "unlink":
            output.unlink()
        return _WorkerRemediationResult(result=result)

    with (
        patch(
            "src.api.brightspace_routes._run_brightspace_worker",
            new=run_worker,
        ),
        patch(
            "src.api.brightspace_routes.RemediationArtifactService.from_settings",
            return_value=service,
        ),
        patch(
            "src.education.validation.matterhorn.MatterhornValidator",
            return_value=validator,
        ),
        patch(
            "src.services.scan_fix_service.lock_scan_review_graph",
            return_value=ScanReviewGraph(
                SimpleNamespace(id="scan-pdf", current_remediation_artifact_id=None),
                (),
                (),
                (),
            ),
        ),
    ):
        outcome = await _remediate_file_impl(
            cloud_file,
            db,
            api_client=api_client,
        )
    return SimpleNamespace(
        outcome=outcome,
        cloud_file=cloud_file,
        db=db,
        service=service,
        published=published,
        validated=validated,
        artifact=artifacts[-1] if artifacts else None,
        scan=db.source_scan,
    )


def _image_equation_result(output_path: Path) -> RemediationResult:
    output_path.write_bytes(CLAIMED_PDF)
    fixed = _typed_equation_fix(
        issue_id="image-equation-1",
        description="Image equation lacked an accessible formula",
        location="page 1 / image 0 / occurrence 0",
        fixed_content="Associated verified Formula, Alt, and MathML",
        provider_used="ollama",
        model_used="vision-test",
        page_number=1,
    )
    result = RemediationResult(
        original_file=str(output_path.with_name("source.pdf")),
        output_file=str(output_path),
        document_type="PDF",
        total_issues=1,
        fixed_count=1,
        manual_count=0,
        failed_count=0,
        fixed_issues=[fixed],
        verification_passed=True,
        success=True,
    )
    result.set_output_claim(
        DescriptorBoundOutputClaim._snapshot_from_owned_descriptor(
            os.open(output_path, os.O_RDONLY),
            filename=output_path.name,
            display_path=str(output_path),
            mime="application/pdf",
        )
    )
    return result


def _real_partial_pdf_result(tmp_path):
    import pymupdf as fitz

    from src.api.brightspace_routes import _run_remediator_worker
    from src.education.pdf_processor import PDFProcessor
    from src.education.remediation.base import RemediationConfig

    source = tmp_path / "source.pdf"
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_textbox(
            fitz.Rect(40, 40, 540, 700),
            "Accessible course material with a clear reading order. " * 30,
        )
        pdf.save(source)
    issues = (
        PDFProcessor(generate_alt_text=False, validate_alt_text=False)
        .process_pdf(str(source))
        .issues
    )
    worker = _run_remediator_worker(
        ext="pdf",
        raw_issues=issues,
        config=RemediationConfig(
            use_ai=False, allow_legacy_nested_ai=False, create_backup=False
        ),
        remediation_client=None,
        source_bytes=source.read_bytes(),
    )
    assert worker.result.verification_passed is True
    assert worker.result.fixed_count > 0
    assert worker.result.manual_count > 0
    assert worker.result.score_measurement is not None
    return source.read_bytes(), worker.result


def _approve_produced_artifact(run, tmp_path, *, change=None):
    from io import BytesIO

    from src.services.remediation_artifact_service import RemediationArtifactService
    from src.services.scan_fix_service import apply_authenticated_batch_review

    rows = [
        call.args[0]
        for call in run.db.add.call_args_list
        if call.args and isinstance(call.args[0], ScanFix)
    ]
    assert rows
    if change == "edit":
        rows[0].fixed_content = "Changed after these output bytes were saved"
    apply_authenticated_batch_review(
        run.db,
        scan_id=run.scan.id,
        fixes=rows,
        action="approve",
        user_id="human-reviewer",
        reviewed_at=datetime.now(timezone.utc),
    )
    if change == "reject":
        apply_authenticated_batch_review(
            run.db,
            scan_id=run.scan.id,
            fixes=rows[:1],
            action="reject",
            user_id="human-reviewer",
            reviewed_at=datetime.now(timezone.utc),
        )
    if change == "output":
        run.artifact.sha256 = "0" * 64
    query = MagicMock()
    query.filter.return_value = query
    query.with_for_update.return_value = query
    query.populate_existing.return_value = query
    query.all.return_value = rows
    run.db.query.side_effect = lambda _: query
    service = RemediationArtifactService(
        root=tmp_path / "managed",
        max_bytes=1_000_000,
        retention_days=1,
        staging_grace_seconds=1,
    )

    @contextmanager
    def verified(_db, artifact, **_authority):
        payload = run.published[0]
        assert hashlib.sha256(payload).hexdigest() == artifact.sha256
        with BytesIO(payload) as stream:
            yield stream

    with (
        patch.object(
            service,
            "_lock_mutable_graph",
            return_value=(run.scan, run.cloud_file, run.artifact),
        ),
        patch.object(service, "open_verified", side_effect=verified),
    ):
        return service.approve(
            run.db,
            artifact_id=run.artifact.id,
            approved_by_id="human-reviewer",
            approved_by_ref="session:human-reviewer",
        )


@pytest.mark.asyncio
async def test_real_partial_pdf_producer_can_reach_human_artifact_approval(tmp_path):
    source, result = _real_partial_pdf_result(tmp_path)
    run = await _run_pdf_outer(tmp_path, result, source_bytes=source)
    assert run.outcome.status == "manual_required"
    assert run.scan.remediation_outcome == "completed"
    receipt = run.artifact.provider_result["reviewed_output_membership"]
    assert receipt["source_sha256"] == hashlib.sha256(source).hexdigest()
    assert receipt["output_sha256"] == hashlib.sha256(run.published[0]).hexdigest()
    assert len(receipt["fixes"]) == result.fixed_count
    approved = _approve_produced_artifact(run, tmp_path)
    assert approved.review_status == "approved"
    assert approved.approved_by_id == "human-reviewer"
    assert run.cloud_file.remediated_issues_remaining == result.manual_count


@pytest.mark.asyncio
async def test_real_office_producer_records_applied_members_and_human_approval(
    tmp_path,
):
    from docx import Document

    from src.api.brightspace_routes import _remediate_file_impl
    from src.db.models import ScanResult
    from src.education.docx_processor import DocxProcessor

    source = tmp_path / "source.docx"
    document = Document()
    document.add_heading("Course overview", level=1)
    document.add_paragraph("Review this week's notes.")
    document.save(source)
    source_bytes = source.read_bytes()
    issues = DocxProcessor(require_complete_scan=True).process_docx(str(source)).issues
    cloud = _cloud_file()
    cloud.file_name = "source.docx"
    cloud.provider_metadata["url"] = "/content/source.docx"
    cloud.file_size_bytes = len(source_bytes)
    db = _db()
    db.source_scan.scan_type = "WORD"
    db.source_scan.file_hash = hashlib.sha256(source_bytes).hexdigest()
    db.query(ScanResult).first.return_value.issues = issues
    api = AsyncMock()
    api.get_topic_file.return_value = (source_bytes, "application/octet-stream")
    artifact = _artifact()
    published = []
    service = MagicMock()

    def publish(_db, **kwargs):
        payload = Path(kwargs["source_path"]).read_bytes()
        published.append(payload)
        artifact.sha256 = hashlib.sha256(payload).hexdigest()
        artifact.size_bytes = len(payload)
        artifact.provider_result = kwargs["provider_result"]
        cloud.current_remediation_artifact_id = artifact.id
        return ArtifactPublicationResult(artifact=artifact, artifact_id=artifact.id)

    service.claim_and_publish.side_effect = publish

    async def run_worker(_department, worker, *args, **kwargs):
        return worker(*args, **kwargs)

    with (
        patch("src.api.brightspace_routes._run_brightspace_worker", new=run_worker),
        patch(
            "src.api.brightspace_routes.RemediationArtifactService.from_settings",
            return_value=service,
        ),
        patch(
            "src.services.scan_fix_service.lock_scan_review_graph",
            return_value=ScanReviewGraph(db.source_scan, (), (), ()),
        ),
    ):
        outcome = await _remediate_file_impl(cloud, db, api_client=api)
    assert outcome.status == "completed"
    assert db.source_scan.remediation_outcome == "completed"
    assert (
        artifact.provider_result["reviewed_output_membership"]["source_sha256"]
        == db.source_scan.file_hash
    )
    run = SimpleNamespace(
        db=db,
        scan=db.source_scan,
        artifact=artifact,
        cloud_file=cloud,
        published=published,
    )
    assert _approve_produced_artifact(run, tmp_path).review_status == "approved"


@pytest.mark.asyncio
@pytest.mark.parametrize("change", ["edit", "reject", "output"])
async def test_review_changes_cannot_approve_old_brightspace_output(tmp_path, change):
    from src.services.remediation_artifact_service import ArtifactAuthorizationError

    source, result = _real_partial_pdf_result(tmp_path)
    run = await _run_pdf_outer(tmp_path, result, source_bytes=source)
    with pytest.raises(ArtifactAuthorizationError):
        _approve_produced_artifact(run, tmp_path, change=change)
    assert run.artifact.approved_by_id is None


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "change,reason",
    [
        ("missing_source", "source_hash_unavailable"),
        ("stale_source", "source_changed_since_scan"),
        ("concurrent_source", "source_changed_since_scan"),
        ("missing_verification", "saved_file_verification_unavailable"),
        ("wrong_output_receipt", "saved_file_verification_unavailable"),
    ],
)
async def test_unbound_partial_output_stays_downloadable_but_cannot_be_approved(
    tmp_path, change, reason
):
    from src.services.remediation_artifact_service import ArtifactAuthorizationError

    source, result = _real_partial_pdf_result(tmp_path)
    options = {}
    if change == "missing_source":
        options["source_hash"] = None
    elif change == "stale_source":
        options["source_hash"] = "0" * 64
    elif change == "concurrent_source":
        options["final_source_hash"] = "0" * 64
    elif change == "missing_verification":
        result.verification_result = None
    else:
        result.score_measurement["output_sha256"] = "0" * 64
    run = await _run_pdf_outer(tmp_path, result, source_bytes=source, **options)
    assert run.outcome.has_remediated_version is True
    assert run.published
    assert run.scan.remediation_outcome is None
    assert "reviewed_output_membership" not in run.artifact.provider_result
    assert run.artifact.provider_result["output_membership_blocker"] == reason
    with pytest.raises(ArtifactAuthorizationError):
        _approve_produced_artifact(run, tmp_path)


@pytest.mark.asyncio
async def test_brightspace_pdf_persists_real_image_equation_review_evidence(
    tmp_path,
):
    result = _image_equation_result(tmp_path / "fixed.pdf")

    run = await _run_pdf_outer(tmp_path, result)

    assert run.outcome.status == "completed"
    persisted = [
        call.args[0]
        for call in run.db.add.call_args_list
        if call.args and isinstance(call.args[0], ScanFix)
    ]
    assert len(persisted) == 1
    row = persisted[0]
    assert row.issue_id == "image-equation-1"
    assert row.source_kind == "image_equation"
    assert row.provider_used == "ollama"
    assert row.model_used == "vision-test"
    assert row.fix_method == "ai_vision"
    assert row.confidence == 0.55
    assert row.needs_review is True
    assert row.review_status == "pending"
    assert row.verification_evidence["threshold_version"] == "printed-equation-v1"
    assert row.verification_evidence["source_sha256"] == EVIDENCE["source_sha256"]


@pytest.mark.asyncio
@pytest.mark.parametrize(
    ("matterhorn_result", "matterhorn_error"),
    [
        (None, RuntimeError("validator unavailable")),
        (
            SimpleNamespace(checkpoints=[], total=0, passed=0, failed=0, warnings=0),
            None,
        ),
        (
            SimpleNamespace(
                checkpoints=[SimpleNamespace(status=SimpleNamespace(value="fail"))],
                total=1,
                passed=0,
                failed=1,
                warnings=0,
            ),
            None,
        ),
        (
            SimpleNamespace(
                checkpoints=[SimpleNamespace(status=SimpleNamespace(value="pass"))],
                total=2,
                passed=1,
                failed=0,
                warnings=0,
            ),
            None,
        ),
    ],
    ids=["exception", "unavailable", "disqualifying", "integrity"],
)
async def test_brightspace_image_equation_matterhorn_failure_preserves_prior_state(
    tmp_path, matterhorn_result, matterhorn_error
):
    result = _image_equation_result(tmp_path / "fixed.pdf")

    run = await _run_pdf_outer(
        tmp_path,
        result,
        matterhorn_result=matterhorn_result,
        matterhorn_error=matterhorn_error,
    )

    assert run.outcome.status == "failed"
    assert run.outcome.error_code == "artifact_unavailable"
    assert run.published == []
    assert run.cloud_file.has_remediated_version is False
    assert run.cloud_file.remediation_origin is None
    assert run.cloud_file.remediated_issues_fixed == 0
    assert run.cloud_file.remediated_issues_remaining == 0
    assert result.has_output_claim() is False


@pytest.mark.asyncio
@pytest.mark.parametrize("tamper", ["replace", "truncate", "unlink"])
async def test_brightspace_pdf_publishes_and_validates_exact_claim_after_path_tamper(
    tmp_path, tamper
):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)

    run = await _run_pdf_outer(tmp_path, result, tamper=tamper)

    assert run.outcome.status == "completed"
    assert run.outcome.artifact_sha256 == hashlib.sha256(CLAIMED_PDF).hexdigest()
    assert run.published == [CLAIMED_PDF]
    assert run.validated == [CLAIMED_PDF]
    run.service.claim_and_publish.assert_not_called()
    assert result.close_calls == 1
    assert result.has_output_claim() is False
    serialized = run.outcome.model_dump()
    assert all("path" not in key and "fd" not in key for key in serialized)


def test_brightspace_pdf_worker_never_reopens_claimed_output_path(monkeypatch):
    from src.api.brightspace_routes import _run_remediator_worker

    state: dict[str, object] = {"remediation_returned": False}
    real_open = builtins.open
    real_isfile = os.path.isfile

    class FakePdfRemediator:
        def __init__(self, file_path, *_args, **_kwargs):
            self.file_path = Path(file_path)

        def remediate(self):
            output = self.file_path.with_name("fixed.pdf")
            output.write_bytes(CLAIMED_PDF)
            result = _PdfResult(output)
            state["output"] = output
            state["result"] = result
            state["remediation_returned"] = True
            return result

    def guarded_open(file, *args, **kwargs):
        if state["remediation_returned"] and Path(file) == state["output"]:
            raise AssertionError("PDF output pathname was reopened")
        return real_open(file, *args, **kwargs)

    def guarded_isfile(path):
        if state["remediation_returned"] and Path(path) == state["output"]:
            raise AssertionError("PDF output pathname was restatted")
        return real_isfile(path)

    module = SimpleNamespace(PdfRemediator=FakePdfRemediator)
    monkeypatch.setattr(builtins, "open", guarded_open)
    monkeypatch.setattr(os.path, "isfile", guarded_isfile)
    monkeypatch.setattr(
        "src.api.brightspace_routes.importlib.import_module",
        lambda *_args, **_kwargs: module,
    )

    worker_result = _run_remediator_worker(
        ext="pdf",
        raw_issues=[{"category": "structure"}],
        config=SimpleNamespace(),
        remediation_client=None,
        source_bytes=b"%PDF-source",
    )
    result = worker_result.result
    try:
        assert worker_result.remediated_bytes is None
        assert result.has_output_claim() is True
        with result.open_output_stream() as stream:
            assert stream.read() == CLAIMED_PDF
        assert Path(result.output_file).exists() is False
    finally:
        result.close_output_claim()


def test_brightspace_non_pdf_worker_keeps_remediated_bytes(monkeypatch):
    from src.api.brightspace_routes import _run_remediator_worker

    office_bytes = b"PK\x03\x04remediated-docx"

    class FakeDocxRemediator:
        def __init__(self, file_path, *_args, **_kwargs):
            self.file_path = Path(file_path)

        def remediate(self):
            output = self.file_path.with_name("fixed.docx")
            output.write_bytes(office_bytes)
            return SimpleNamespace(
                success=True,
                output_file=str(output),
                verification_passed=True,
                fixed_count=1,
                manual_count=0,
                failed_count=0,
            )

    monkeypatch.setattr(
        "src.api.brightspace_routes.importlib.import_module",
        lambda *_args, **_kwargs: SimpleNamespace(DocxRemediator=FakeDocxRemediator),
    )

    worker_result = _run_remediator_worker(
        ext="docx",
        raw_issues=[{"category": "heading"}],
        config=SimpleNamespace(),
        remediation_client=None,
        source_bytes=b"PK\x03\x04source-docx",
    )

    assert worker_result.remediated_bytes == office_bytes


@pytest.mark.asyncio
@pytest.mark.parametrize("claim_state", ["missing", "closed"])
async def test_brightspace_pdf_missing_or_closed_claim_is_artifact_unavailable(
    tmp_path, claim_state
):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output, with_claim=claim_state == "closed")
    if result.claim is not None:
        result.claim.close()

    run = await _run_pdf_outer(tmp_path, result)

    assert run.outcome.status == "failed"
    assert run.outcome.error_code == "artifact_unavailable"
    assert run.outcome.has_remediated_version is False
    run.service.claim_and_publish_stream.assert_not_called()
    assert run.validated == []
    assert result.close_calls == 1


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "result_kwargs",
    [
        {"metadata_error": RuntimeError("metadata race")},
        {
            "metadata_override": {
                "size_bytes": len(CLAIMED_PDF),
                "sha256": "not-a-digest",
                "mime_type": "application/pdf",
                "filename": "fixed.pdf",
            }
        },
        {"fail_stream_on": 2},
    ],
)
async def test_brightspace_pdf_metadata_and_stream_races_fail_closed(
    tmp_path, result_kwargs
):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output, **result_kwargs)

    run = await _run_pdf_outer(tmp_path, result)

    assert run.outcome.status == "failed"
    assert run.outcome.error_code == "artifact_unavailable"
    assert run.outcome.has_remediated_version is False
    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_pdf_publication_failure_closes_claim_and_fails_closed(
    tmp_path,
):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)

    run = await _run_pdf_outer(
        tmp_path,
        result,
        publication_error=OSError("artifact store unavailable"),
    )

    assert run.outcome.status == "failed"
    assert run.outcome.error_code == "artifact_unavailable"
    assert run.outcome.has_remediated_version is False
    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_pdf_db_failure_closes_claim(tmp_path):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)

    run = await _run_pdf_outer(
        tmp_path,
        result,
        commit_error=RuntimeError("database unavailable"),
    )

    assert run.outcome.status == "failed"
    assert run.outcome.error_code == "remediation_failed"
    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_worker_cancellation_waits_then_closes_returned_claim(
    tmp_path,
):
    from src.api.brightspace_routes import (
        _WorkerRemediationResult,
        _run_brightspace_worker,
    )

    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    started = threading.Event()
    release = threading.Event()

    def worker():
        started.set()
        assert release.wait(timeout=5)
        return _WorkerRemediationResult(result=result)

    task = asyncio.create_task(_run_brightspace_worker("dept-cancel", worker))
    assert await asyncio.to_thread(started.wait, 2)
    task.cancel()
    await asyncio.sleep(0)
    assert task.done() is False
    assert result.has_output_claim() is True
    release.set()
    with pytest.raises(asyncio.CancelledError):
        await task

    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_pdf_publish_cancellation_closes_claim(tmp_path):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)

    with pytest.raises(asyncio.CancelledError):
        await _run_pdf_outer(
            tmp_path,
            result,
            publication_error=asyncio.CancelledError(),
        )

    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_pdf_real_cancellation_during_publish_is_observed(tmp_path):
    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    publish_started = threading.Event()
    publish_release = threading.Event()
    loop = asyncio.get_running_loop()
    task = None

    def block_publish():
        publish_started.set()
        assert publish_release.wait(timeout=5)

    def cancel_during_publish():
        assert publish_started.wait(timeout=5)
        assert task is not None
        loop.call_soon_threadsafe(task.cancel)
        publish_release.set()

    controller = threading.Thread(target=cancel_during_publish)
    controller.start()
    task = asyncio.create_task(
        _run_pdf_outer(
            tmp_path,
            result,
            publication_hook=block_publish,
        )
    )
    with pytest.raises(asyncio.CancelledError):
        await task
    controller.join(timeout=5)

    assert controller.is_alive() is False
    assert result.close_calls == 1
    assert result.has_output_claim() is False


@pytest.mark.asyncio
async def test_brightspace_pdf_cancellation_after_sync_publish_aborts_exact_claim(
    tmp_path,
):
    from src.api.brightspace_routes import _finish_brightspace_pdf_remediation

    output = tmp_path / "fixed.pdf"
    output.write_bytes(CLAIMED_PDF)
    result = _PdfResult(output)
    cloud_file = _cloud_file()
    db = _db()
    artifact = _artifact()
    publication = ArtifactPublicationResult(
        artifact=artifact,
        artifact_id=str(artifact.id),
        publication_token="-".join(("brightspace", "publication", "fixture")),
    )
    service = MagicMock()
    task = asyncio.current_task()
    loop = asyncio.get_running_loop()

    def publish(*args, **kwargs):
        cloud_file.has_remediated_version = True
        cloud_file.remediation_origin = "manual"
        loop.call_soon(task.cancel)
        return publication

    service.claim_and_publish_stream.side_effect = publish
    with (
        patch(
            "src.api.brightspace_routes.RemediationArtifactService.from_settings",
            return_value=service,
        ),
        patch(
            "src.api.brightspace_routes._validate_brightspace_pdf_claim",
        ),
    ):
        with pytest.raises(asyncio.CancelledError):
            await _finish_brightspace_pdf_remediation(
                cloud_file,
                db,
                result=result,
                complete=True,
                decisions={"remediation": "used", "alt_text": "not_requested"},
                alt_text_client=None,
            )

    service.abort_staging.assert_called_once_with(
        db,
        artifact_id=publication.artifact_id,
        publication_token=publication.publication_token,
    )
    db.rollback.assert_called_once_with()
    assert cloud_file.has_remediated_version is False
    assert cloud_file.remediation_origin is None
    assert cloud_file.writeback_status is None
    assert result.close_calls == 1
    assert result.has_output_claim() is False
