"""Saved Office acceptance evidence must detect loss, not just higher scores."""

from pathlib import Path
import socket
import json
import subprocess
import sys
from datetime import datetime, timezone
from types import SimpleNamespace
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
import pytest
import yaml

from scripts import office_preservation_corpus as corpus
from scripts.office_preservation_corpus import (
    compare_packages,
    generate_sources,
    run_corpus,
)
from scripts.pdf_acceptance_corpus import sha256_file


def test_office_sources_are_reproducible(tmp_path, monkeypatch):
    from openpyxl.writer import excel

    times = iter((datetime(2001, 1, 1), datetime(2002, 2, 2)))
    monkeypatch.setattr(
        excel,
        "datetime",
        SimpleNamespace(
            datetime=SimpleNamespace(now=lambda **kwargs: next(times)),
            timezone=timezone,
        ),
    )
    first = generate_sources(tmp_path / "first")
    second = generate_sources(tmp_path / "second")
    assert set(first) == set(second) == {"docx", "pptx", "xlsx"}
    assert {k: sha256_file(p) for k, p in first.items()} == {
        k: sha256_file(p) for k, p in second.items()
    }


def test_real_saved_office_corpus(tmp_path, monkeypatch):
    attempts = []

    def no_network(*args, **kwargs):
        attempts.append(True)
        raise AssertionError("External connections are forbidden in the corpus")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    monkeypatch.setattr(socket.socket, "connect_ex", no_network)
    report = run_corpus(tmp_path)
    assert attempts == []
    assert report["status"] == "passed", report
    assert report["failed"] == 0
    assert report["passed"] == 6
    assert report["claim_boundary"]
    for case in report["cases"]:
        assert case["source_sha256_before"] == case["source_sha256_after"]
        assert case["preservation_failures"] == []
        assert case["candidate_sha256"] == sha256_file(tmp_path / case["candidate"])
        assert not Path(case["candidate"]).is_absolute()


def rewrite_package(source, destination, part, transform):
    with ZipFile(source) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    assert part in members
    previous = members[part]
    members[part] = transform(previous)
    assert previous != members[part], "The corruption probe must actually change bytes"
    with ZipFile(destination, "w", ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)


def alter_xml(payload, local_name, *, attribute=None):
    root = etree.fromstring(payload)
    nodes = root.xpath(f"//*[local-name()='{local_name}']")
    assert nodes, "The fixture must contain the property under test"
    if attribute:
        nodes[0].set(attribute, "https://example.org/changed")
    else:
        nodes[0].text = "CORRUPTED"
    return etree.tostring(root)


MUTATIONS = [
    ("docx", "word/document.xml", "t", None),
    ("docx", "word/header1.xml", "t", None),
    ("docx", "word/footer1.xml", "t", None),
    ("docx", "word/_rels/document.xml.rels", "Relationship", "Target"),
    ("pptx", "ppt/slides/slide1.xml", "t", None),
    ("pptx", "ppt/notesSlides/notesSlide1.xml", "t", None),
    ("pptx", "ppt/slides/slide2.xml", "cTn", "dur"),
    ("xlsx", "xl/worksheets/sheet1.xml", "f", None),
    ("xlsx", "xl/worksheets/sheet1.xml", "v", None),
    ("xlsx", "xl/workbook.xml", "definedName", None),
    ("xlsx", "xl/charts/chart1.xml", "f", None),
    ("xlsx", "xl/comments/comment1.xml", "t", None),
]


@pytest.mark.parametrize("kind,part,tag,attribute", MUTATIONS)
def test_package_oracle_rejects_non_target_changes(
    tmp_path, kind, part, tag, attribute
):
    source = generate_sources(tmp_path / "source")[kind]
    assert compare_packages(source, source, kind, repair=True) == []
    candidate = tmp_path / f"candidate.{kind}"
    rewrite_package(
        source, candidate, part, lambda p: alter_xml(p, tag, attribute=attribute)
    )
    assert f"changed:{part}" in compare_packages(source, candidate, kind, repair=True)


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_readable_candidate_can_still_fail_preservation(tmp_path, kind):
    source = generate_sources(tmp_path / "source")[kind]
    part, tag = {
        "docx": ("word/document.xml", "t"),
        "pptx": ("ppt/slides/slide1.xml", "t"),
        "xlsx": ("xl/worksheets/sheet1.xml", "f"),
    }[kind]
    candidate = tmp_path / f"candidate.{kind}"
    rewrite_package(source, candidate, part, lambda p: alter_xml(p, tag))
    # Independent ordinary readers still accept the damaged packages.
    if kind == "docx":
        assert Document(candidate).paragraphs
    elif kind == "pptx":
        assert len(Presentation(candidate).slides) == 2
    else:
        workbook = load_workbook(candidate)
        assert workbook.active["C2"].value == "=CORRUPTED"
        workbook.close()
    assert compare_packages(source, candidate, kind, repair=True)


@pytest.mark.parametrize(
    "damage", ["media", "relationship", "shape-order", "missing-part", "added-part"]
)
def test_package_inventory_and_presentation_relationships_are_preserved(
    tmp_path, damage
):
    source = generate_sources(tmp_path / "source")["pptx"]
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    if damage == "media":
        name = next(
            name for name, payload in parts.items() if payload.startswith(b"RIFF")
        )
        parts[name] += b"changed"
    elif damage == "relationship":
        name = "ppt/slides/_rels/slide1.xml.rels"
        parts[name] = alter_xml(parts[name], "Relationship", attribute="Target")
    elif damage == "shape-order":
        name = "ppt/slides/slide1.xml"
        root = etree.fromstring(parts[name])
        tree = root.find("p:cSld/p:spTree", corpus.NS)
        tree.append(tree[2])
        parts[name] = etree.tostring(root)
    elif damage == "missing-part":
        parts.pop("ppt/notesSlides/notesSlide1.xml")
    else:
        parts["extra.txt"] = b"unexpected content"
    candidate = tmp_path / "changed.pptx"
    with ZipFile(candidate, "w", ZIP_DEFLATED) as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    assert compare_packages(source, candidate, "pptx", repair=True)


@pytest.mark.parametrize("damage", ["text", "wrong-title"])
def test_runner_rejects_corrupt_saved_output_even_after_successful_scan(
    tmp_path, monkeypatch, damage
):
    from src.education.remediation.docx_remediator import DocxRemediator

    save = DocxRemediator._save_document

    def corrupted_save(self, document):
        if damage == "text":
            document.paragraphs[1].text = "Unintended content replacement"
        else:
            document.core_properties.title = "Wrong title"
        return save(self, document)

    monkeypatch.setattr(DocxRemediator, "_save_document", corrupted_save)
    report = run_corpus(tmp_path)
    assert report["status"] == "failed"
    case = next(c for c in report["cases"] if c["id"] == "docx-repair")
    assert case["verified_fixes"] == 1
    expected = (
        "changed:word/document.xml" if damage == "text" else "target_not_persisted"
    )
    assert expected in case["preservation_failures"]


def test_manifest_inventory_cannot_silently_lose_a_case(tmp_path, monkeypatch):
    value = json.loads(corpus.MANIFEST.read_text())
    value["required_cases"].pop()
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps(value))
    monkeypatch.setattr(corpus, "MANIFEST", manifest)
    with pytest.raises(ValueError, match="contract drift"):
        run_corpus(tmp_path / "run")


def test_runner_checks_masked_freeze_pane_attributes(tmp_path, monkeypatch):
    from src.education.remediation.xlsx_remediator import XlsxRemediator

    save = XlsxRemediator._save_document

    def corrupted_save(self, workbook):
        workbook["Course results"].sheet_view.pane.ySplit = 2
        return save(self, workbook)

    monkeypatch.setattr(XlsxRemediator, "_save_document", corrupted_save)
    report = run_corpus(tmp_path)
    case = next(c for c in report["cases"] if c["id"] == "xlsx-repair")
    assert case["verified_fixes"] == 1
    assert "target_not_persisted" in case["preservation_failures"]


def test_existing_corpus_sources_are_never_overwritten(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    sentinel = source / "keep.txt"
    sentinel.write_text("Operator content")
    with pytest.raises(FileExistsError):
        run_corpus(tmp_path)
    assert sentinel.read_text() == "Operator content"


def test_unavailable_evidence_cannot_be_promoted_to_pass(tmp_path, monkeypatch):
    value = json.loads(corpus.MANIFEST.read_text())
    value["unavailable_evidence"][0]["status"] = "passed"
    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps(value))
    monkeypatch.setattr(corpus, "MANIFEST", manifest)
    with pytest.raises(ValueError, match="cannot count as passing"):
        run_corpus(tmp_path / "run")


def test_cli_emits_revision_bound_evidence_and_refuses_reuse(tmp_path):
    workdir = tmp_path / "evidence"
    command = [
        sys.executable,
        str(corpus.ROOT / "scripts/office_preservation_corpus.py"),
        "--workdir",
        str(workdir),
        "--revision",
        "0" * 40,
    ]
    result = subprocess.run(command, capture_output=True, text=True, timeout=60)
    assert result.returncode == 0, result.stderr
    report = json.loads((workdir / "report.json").read_text())
    assert report["status"] == "passed"
    assert report["revision"] == "0" * 40
    assert report["manifest_sha256"] == sha256_file(corpus.MANIFEST)
    assert report["configuration"]["use_ai"] is False
    assert all(report["dependency_versions"].values())
    original = (workdir / "report.json").read_bytes()
    rerun = subprocess.run(command, capture_output=True, text=True, timeout=60)
    assert rerun.returncode != 0
    assert (workdir / "report.json").read_bytes() == original


def test_ci_runs_corpus_and_retains_saved_evidence():
    workflow = yaml.safe_load((corpus.ROOT / ".github/workflows/ci.yml").read_text())
    steps = workflow["jobs"]["test"]["steps"]
    run = next(
        step
        for step in steps
        if step.get("name") == "Required Office saved-package preservation corpus"
    )
    assert not run.get("continue-on-error", False)
    assert '--revision "${{ github.sha }}"' in run["run"]
    assert "scripts/office_preservation_corpus.py" in run["run"]
    archive = next(
        step
        for step in steps
        if step.get("name") == "Retain Office source and candidate evidence"
    )
    assert archive["if"] == "always()"
    assert archive["with"]["path"] == "test-results/office-preservation/"
