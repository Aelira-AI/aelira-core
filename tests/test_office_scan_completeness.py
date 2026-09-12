"""A swallowed Office checker error cannot become a healthy numeric score."""

import shutil
import zipfile
from xml.etree import ElementTree as ET

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from src.education.docx_processor import DocxProcessor
from src.education.pptx_processor import PowerPointProcessor
from src.education.xlsx_processor import XlsxProcessor
from src.education.remediation.base import RemediationConfig
from src.education.remediation.docx_remediator import DocxRemediator
from src.education.remediation.pptx_remediator import PptxRemediator
from src.education.remediation.xlsx_remediator import XlsxRemediator
from src.education.scan_completeness import (
    IncompleteScanError,
    complete_scan_requested,
    require_complete_scan,
)


def make_document(tmp_path, kind):
    source = tmp_path / f"source.{kind}"
    if kind == "docx":
        document = Document()
        document.add_heading("Course", level=2)
        document.add_paragraph("Course materials")
        document.save(source)
        # SmartArt is discovered independently from the package's diagram parts.
        with zipfile.ZipFile(source, "a") as package:
            package.writestr(
                "word/diagrams/data1.xml", "<diagram><t>Course flow</t></diagram>"
            )
        scan = DocxProcessor().process_docx
        remediator = DocxRemediator
    elif kind == "pptx":
        document = Presentation()
        document.slides.add_slide(document.slide_layouts[0])
        document.save(source)
        scan = PowerPointProcessor().process_pptx
        remediator = PptxRemediator
    else:
        document = Workbook()
        document.active.append(["Course", "Hours"])
        document.active.append(["History", 3])
        document.save(source)
        scan = XlsxProcessor().process_xlsx
        remediator = XlsxRemediator
    return source, scan, remediator


def fail_checker_for_path(monkeypatch, kind, failing_path):
    if kind in {"docx", "pptx"}:
        original = ET.parse
        target = (
            "word/diagrams/data1.xml"
            if kind == "docx"
            else "ppt/slides/_rels/slide1.xml.rels"
        )

        def broken_xml(stream, *args, **kwargs):
            archive = getattr(getattr(stream, "_fileobj", None), "_file", None)
            if getattr(stream, "name", None) == target and getattr(
                archive, "name", None
            ) == str(failing_path):
                raise ET.ParseError("Injected analysis XML failure")
            return original(stream, *args, **kwargs)

        monkeypatch.setattr(ET, "parse", broken_xml)
    else:
        process = XlsxProcessor._process_xlsx
        color = XlsxProcessor._get_font_color_hex

        def remember_path(self, path, *args, **kwargs):
            self._fault_fixture_path = path
            return process(self, path, *args, **kwargs)

        def invalid_output_color(self, cell):
            if self._fault_fixture_path == str(failing_path):
                return "#invalid"
            return color(self, cell)

        monkeypatch.setattr(XlsxProcessor, "_process_xlsx", remember_path)
        monkeypatch.setattr(XlsxProcessor, "_get_font_color_hex", invalid_output_color)


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_source_scan_refuses_number_when_required_checker_fails(
    tmp_path, monkeypatch, kind
):
    source, scan, _ = make_document(tmp_path, kind)
    control = scan(str(source))
    assert control.issues
    fail_checker_for_path(monkeypatch, kind, source)
    with pytest.raises(IncompleteScanError):
        scan(str(source))
    assert not complete_scan_requested()


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_swallowed_output_checker_failure_withholds_score(tmp_path, monkeypatch, kind):
    source, scan, remediator_type = make_document(tmp_path, kind)
    before = scan(str(source))
    output = tmp_path / f"output.{kind}"
    shutil.copyfile(source, output)
    config = RemediationConfig(create_backup=False)
    control = remediator_type(str(source), before.issues, config)
    control._verify_fixes(str(output))
    assert control.result.remediated_compliance_score == before.compliance_score
    assert control.result.score_provenance == "scanner_rescan"

    fail_checker_for_path(monkeypatch, kind, output)
    assert scan(str(source)).compliance_score == before.compliance_score
    remediator = remediator_type(str(source), before.issues, config)
    remediator._verify_fixes(str(output))
    result = remediator.result
    assert result.original_compliance_score == before.compliance_score
    assert result.remediated_compliance_score is None
    assert result.score_provenance is None
    assert result.fixed_count == 0
    assert not result.verification_passed


def test_xlsx_failed_contrast_calculation_cannot_claim_maximum_contrast():
    scanner = XlsxProcessor()
    with require_complete_scan(True):
        assert scanner._calculate_contrast_ratio("#000000", "#ffffff") == 21
    with pytest.raises(IncompleteScanError, match="contrast"):
        with require_complete_scan(True):
            scanner._calculate_contrast_ratio("#invalid", "#ffffff")
