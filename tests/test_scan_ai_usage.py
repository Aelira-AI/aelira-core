"""Usage must describe observed inference, never requested options."""

from types import SimpleNamespace
from unittest.mock import MagicMock
import asyncio

import pytest
from docx import Document
from docx.shared import Inches

from test_workspace_provider_runtime import _row, _runtime, _snapshot


@pytest.mark.parametrize("mode", ["disabled", "failed", "cloud", "ollama"])
def test_word_persistence_measures_successful_local_calls(tmp_path, monkeypatch, mode):
    from pathlib import Path
    from src.ai import workspace_provider_runtime
    from src.api.education import scan_routes
    from src.db import database
    from src.middleware import quota

    primary = None if mode == "disabled" else "openai" if mode == "cloud" else "ollama"
    events = []
    runtime = _runtime(
        lambda _: _snapshot(
            primary=primary,
            rows=[_row(primary, credential="synthetic-cipher")] if primary else [],
        ),
        events=events,
        outcomes={"ollama": mode != "failed"},
    )
    document = Document()
    document.add_heading("Diagram lesson", level=1)
    document.add_picture(
        str(Path(__file__).parent / "fixtures/sample_diagram.png"), width=Inches(4)
    )
    original = tmp_path / "source.docx"
    document.save(original)
    uploaded = tmp_path / "upload.docx"
    uploaded.write_bytes(original.read_bytes())
    scan = SimpleNamespace(id="scan-test", department_id="workspace-a")
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = scan
    monkeypatch.setattr(database, "SessionLocal", lambda: db)
    monkeypatch.setattr(
        workspace_provider_runtime, "workspace_provider_runtime", lambda _: runtime
    )
    monkeypatch.setattr(quota, "increment_usage_sync", lambda *a, **k: None)

    scan_routes.process_docx_background(
        str(uploaded),
        original.read_bytes(),
        original.name,
        scan.id,
        True,
        False,
        str(original),
        "user-a",
        "workspace-a",
    )

    assert db.add.call_count == 1, scan.__dict__
    stored = db.add.call_args.args[0]
    assert stored.ollama_used is (mode == "ollama")
    if mode == "ollama":
        calls = sum(event[0] == "analyze_image" for event in events)
        assert calls > 0
        assert stored.ollama_calls == calls
    else:
        assert stored.ollama_calls == 0


@pytest.mark.asyncio
async def test_runtime_usage_is_per_instance_and_provider_with_fallback():
    snapshot = _snapshot(
        primary="openai",
        fallback="ollama",
        rows=[_row("openai", credential="cipher"), _row("ollama")],
    )
    runtime = _runtime(lambda _: snapshot, outcomes={"openai": False})
    other = _runtime(lambda _: snapshot)
    await runtime.generate_text(prompt="one")
    await runtime.analyze_image(image_data=b"image", prompt="two")
    assert runtime.ollama_usage.successful_calls == 2
    assert other.ollama_usage.successful_calls == 0
    await other.generate_text(prompt="cloud response")
    assert other.ollama_usage.successful_calls == 0


@pytest.mark.asyncio
async def test_concurrent_operations_count_once_and_snapshot_stays_immutable():
    runtime = _runtime(lambda _: _snapshot(primary="ollama", rows=[_row("ollama")]))
    before = runtime.ollama_usage
    await asyncio.gather(*(runtime.generate_text(prompt=str(i)) for i in range(20)))
    assert before.successful_calls == 0
    assert runtime.ollama_usage.successful_calls == 20


def test_unknown_usage_survives_database_round_trip():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session
    from src.db.models import ScanResult

    engine = create_engine("sqlite:///:memory:")
    ScanResult.__table__.create(engine)
    with Session(engine) as db:
        rows = [
            ScanResult(scan_id="unknown", compliance_score=50),
            ScanResult(
                scan_id="explicit-unknown",
                compliance_score=50,
                ollama_used=None,
                ollama_calls=None,
            ),
            ScanResult(
                scan_id="unused", compliance_score=50, ollama_used=False, ollama_calls=0
            ),
        ]
        db.add_all(rows)
        db.commit()
        for row in rows[:2]:
            db.refresh(row)
            assert row.ollama_used is None
            assert row.ollama_calls is None
        db.refresh(rows[2])
        assert rows[2].ollama_used is False
        assert rows[2].ollama_calls == 0
    engine.dispose()


@pytest.mark.parametrize("kind", ["pdf", "powerpoint", "latex", "word", "excel"])
@pytest.mark.parametrize("count", [None, 0, 3])
def test_legacy_storage_requires_measurement_not_suggested_content(kind, count):
    from src.ai.usage import OllamaUsage
    from src.db.models import ScanResult
    from src.db.scan_service import ScanService

    # Deliberately include suggestion-shaped content and inflated object counts.
    # Neither constitutes provider evidence.
    result = SimpleNamespace(
        file_name="fixture",
        pages=1,
        compliance_score=50,
        issues=[{"severity": "high", "how_to_fix": "Manually supplied advice"}],
        structure={},
        html_output="",
        ocr_used=False,
        slides=[],
        total_slides=0,
        total_shapes=99,
        total_images=99,
        remediation_suggestions=["Manually supplied advice"],
        equations=[],
        total_equations=99,
        successful_conversions=0,
        failed_conversions=0,
        image_issues=[],
        heading_issues=[],
        table_issues=[],
        list_issues=[],
        link_issues=[],
        font_issues=[],
        color_issues=[],
        total_paragraphs=0,
        total_tables=0,
        total_lists=0,
        total_links=0,
        sheets=[],
        total_sheets=0,
        total_rows=0,
        total_charts=99,
    )
    db = MagicMock()
    kwargs = {"usage": OllamaUsage(count)} if count is not None else {}
    if kind == "latex":
        from src.education.latex_processor import DocumentConversionResult

        result = DocumentConversionResult(
            file_path="fixture.tex",
            file_name=result.file_name,
            total_equations=result.total_equations,
            successful_conversions=result.successful_conversions,
            failed_conversions=result.failed_conversions,
            equations=result.equations,
            html_output=result.html_output,
            compliance_score=result.compliance_score,
            source_issues=result.issues,
        )
        kwargs["ollama_used"] = True  # Legacy request hint cannot establish use.
    getattr(ScanService, f"store_{kind}_scan")(
        db, result, "user", "workspace", b"fixture", **kwargs
    )
    stored = next(
        call.args[0]
        for call in db.add.call_args_list
        if isinstance(call.args[0], ScanResult)
    )
    assert stored.ollama_used is (None if count is None else count > 0)
    assert stored.ollama_calls == count
