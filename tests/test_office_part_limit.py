"""Large Office files are not refused for their number of parts.

Regression for #874. An Office Open XML file is a zip archive with a part for
each slide, notes page, relationship set and picture, so a deck of about 200
slides passed the 1,000-entry archive limit. The validator rated it HIGH, and
because a rebuild cannot lower that rating the upload was refused. Office
packages now have their own part limit, set from measured decks, and other zip
archives keep the archive limit.
"""

import io
import zipfile
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi import UploadFile
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from src.api.education import _shared
from src.security.document_validator import (
    DocumentValidator,
    ThreatLevel,
)

OFFICE_PART_LIMIT = 10_000
ARCHIVE_FILE_LIMIT = 1_000


@pytest.fixture(scope="session", autouse=True)
def setup_test_database():
    """These tests use an in-memory session and never initialize a database."""
    yield


class MemorySession:
    def __init__(self):
        self.added = []

    def add(self, value):
        self.added.append(value)

    def commit(self):
        pass


@pytest.fixture
def session(monkeypatch):
    monkeypatch.setattr(
        "src.config.settings.get_settings",
        lambda: SimpleNamespace(env="test", max_file_size_pdf=10 * 1024 * 1024),
    )
    return MemorySession()


def picture(seed):
    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), (seed % 256, (seed * 7) % 256, 90)).save(buffer, "PNG")
    return buffer.getvalue()


def deck(slides):
    """A deck with a title, a distinct picture and speaker notes on each slide."""
    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = f"Week {index}"
        slide.shapes.add_picture(io.BytesIO(picture(index)), Inches(1), Inches(2))
        slide.notes_slide.notes_text_frame.text = f"Speaker notes {index}"
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def package(total_parts, *, content_types=True):
    """A zip archive with exactly ``total_parts`` entries."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        written = 0
        if content_types:
            zf.writestr(
                "[Content_Types].xml",
                '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats'
                '.org/package/2006/content-types"></Types>',
            )
            written = 1
        for index in range(total_parts - written):
            zf.writestr(f"word/media/part{index}.xml", "x")
    return buffer.getvalue()


def part_count(content):
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        return len(zf.namelist())


def categories(result):
    return {finding.category for finding in result.findings}


@pytest.mark.asyncio
async def test_large_deck_is_accepted(session):
    content = deck(250)
    assert part_count(content) > ARCHIVE_FILE_LIMIT

    result = await DocumentValidator().validate("lectures.pptx", content)
    assert "archive_file_count" not in categories(result)
    assert result.threat_level.rank < ThreatLevel.HIGH.rank

    accepted = await _shared.validate_uploaded_file(
        UploadFile(filename="lectures.pptx", file=io.BytesIO(content)),
        session,
        "workspace",
    )
    assert accepted == content
    record = session.added[0]
    assert record.was_blocked is False
    assert record.was_sanitized is False


@pytest.mark.asyncio
async def test_office_part_limit_boundary():
    validator = DocumentValidator()

    at_limit = await validator.validate("report.docx", package(OFFICE_PART_LIMIT))
    assert "archive_file_count" not in categories(at_limit)

    past_limit = await validator.validate("report.docx", package(OFFICE_PART_LIMIT + 1))
    finding = next(f for f in past_limit.findings if f.category == "archive_file_count")
    assert finding.threat_level == ThreatLevel.HIGH
    assert finding.description == "Office file contains too many parts (10,001)"
    assert finding.remediation == "Split the document into smaller files"


@pytest.mark.asyncio
async def test_office_package_past_the_part_limit_is_refused(session):
    content = package(OFFICE_PART_LIMIT + 1)

    with pytest.raises(HTTPException) as caught:
        await _shared.validate_uploaded_file(
            UploadFile(filename="report.docx", file=io.BytesIO(content)),
            session,
            "workspace",
        )

    assert caught.value.status_code == 400
    record = session.added[0]
    assert record.was_blocked is True
    assert record.blocked_reason == "Office file contains too many parts (10,001)"


@pytest.mark.asyncio
@pytest.mark.parametrize("filename", ["site.zip", "report.docx.zip"])
async def test_other_archives_keep_the_archive_limit(filename):
    validator = DocumentValidator()

    at_limit = await validator.validate(
        filename, package(ARCHIVE_FILE_LIMIT, content_types=False)
    )
    assert "archive_file_count" not in categories(at_limit)

    past_limit = await validator.validate(
        filename, package(ARCHIVE_FILE_LIMIT + 1, content_types=False)
    )
    finding = next(f for f in past_limit.findings if f.category == "archive_file_count")
    assert finding.threat_level == ThreatLevel.HIGH
    assert finding.description == "Archive contains too many files (1,001)"
