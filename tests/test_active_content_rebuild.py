"""Uploads rated HIGH are rebuilt into readable files without active content.

The packages here are made by python-docx, openpyxl and python-pptx and then
given active content the way Office saves it: a content type and a
relationship for each part, and an element in the document that names it.
"""

import io
import re
import zipfile

import docx
import fitz
import openpyxl
import pikepdf
import pptx
import pytest
from fastapi import HTTPException
from lxml import etree

from src.api.education import _shared
from src.security.active_content import SanitizationError, rebuild_ooxml, rebuild_pdf
from src.security.document_validator import (
    DocumentValidator,
    ThreatLevel,
    sanitize_document,
)
from tests.vba_fixtures import vba_project_bin

pytestmark = pytest.mark.unit

REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
MS_REL = "http://schemas.microsoft.com/office/2006/relationships"
RELS_EMPTY = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"></Relationships>'
)
PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000100e5274de40000000049454e44ae426082"
)
BENIGN_VBA = 'Sub Greeting()\r\n    MsgBox "Hello"\r\nEnd Sub\r\n'


def rel(rid, kind, target, external=False):
    mode = ' TargetMode="External"' if external else ""
    return f'<Relationship Id="{rid}" Type="{kind}" Target="{target}"{mode}/>'


def override(part, content_type):
    return f'<Override PartName="/{part}" ContentType="{content_type}"/>'


def edit_package(base, *, add=None, edit=None, types=(), rels=None, main_type=None):
    with zipfile.ZipFile(io.BytesIO(base)) as zf:
        parts = {info.filename: zf.read(info.filename) for info in zf.infolist()}
    for name, change in (edit or {}).items():
        parts[name] = change(parts[name].decode()).encode()
    parts.update(add or {})
    types_xml = parts["[Content_Types].xml"].decode()
    types_xml = types_xml.replace("</Types>", "".join(types) + "</Types>")
    if main_type:
        types_xml = re.sub(
            r'(PartName="/(?:word/document|xl/workbook|ppt/presentation)\.xml" '
            r'ContentType=")[^"]+',
            rf"\g<1>{main_type}",
            types_xml,
        )
    parts["[Content_Types].xml"] = types_xml.encode()
    for rels_name, entries in (rels or {}).items():
        existing = parts.get(rels_name, RELS_EMPTY.encode()).decode()
        parts[rels_name] = existing.replace(
            "</Relationships>", "".join(entries) + "</Relationships>"
        ).encode()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in parts.items():
            zf.writestr(name, data)
    return buffer.getvalue()


def names(content):
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        return set(zf.namelist())


def read(content, name):
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        return zf.read(name).decode()


def all_relationships(content):
    found = []
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        for name in zf.namelist():
            if name.endswith(".rels"):
                root = etree.fromstring(zf.read(name))
                found += [(name, element.attrib) for element in root]
    return found


def word_base():
    document = docx.Document()
    document.add_heading("Valve practical", level=1)
    document.add_paragraph("Record the reading before you open the valve.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


WORD_OBJECTS = (
    '<w:p xmlns:v="urn:schemas-microsoft-com:vml" '
    'xmlns:o="urn:schemas-microsoft-com:office:office"><w:r>'
    '<w:object w:dxaOrig="1440" w:dyaOrig="1440">'
    '<v:shape id="_x0000_i1025" type="#_x0000_t75" style="width:72pt;height:72pt">'
    '<v:imagedata r:id="rIdImg1" o:title=""/></v:shape>'
    '<o:OLEObject Type="Embed" ProgID="Equation.3" ShapeID="_x0000_i1025" '
    'DrawAspect="Content" ObjectID="_1" r:id="rIdOle1"/></w:object></w:r></w:p>'
    '<w:p xmlns:v="urn:schemas-microsoft-com:vml"><w:r><w:object>'
    '<v:shape id="_x0000_i1026" type="#_x0000_t75">'
    '<v:imagedata r:id="rIdImg1"/></v:shape>'
    '<w:control r:id="rIdAx1" w:name="CheckBox1" w:shapeid="_x0000_i1026"/>'
    "</w:object></w:r></w:p>"
    '<w:p><w:hyperlink r:id="rIdLink1"><w:r><w:t>Aelira website</w:t></w:r>'
    "</w:hyperlink></w:p>"
)


def macro_word_document():
    """A Word file with every kind of active content, saved as .docx."""
    return edit_package(
        word_base(),
        add={
            "word/vbaProject.bin": vba_project_bin(BENIGN_VBA),
            "word/vbaData.xml": "<wne:vbaSuppData xmlns:wne="
            '"http://schemas.microsoft.com/office/word/2006/wordml"/>',
            "word/activeX/activeX1.xml": '<ax:ocx xmlns:ax="http://schemas.microsoft'
            '.com/office/2006/activeX" ax:classid="{8BD21D40-EC42-11CE-9E0D-00AA006002F3}"/>',
            "word/activeX/activeX1.bin": b"\x00activex",
            "word/activeX/_rels/activeX1.xml.rels": RELS_EMPTY.replace(
                "</Relationships>",
                rel("rId1", f"{MS_REL}/activeXControlBinary", "activeX1.bin")
                + "</Relationships>",
            ),
            "word/embeddings/oleObject1.bin": b"\xd0\xcf\x11\xe0 equation",
            "word/media/preview1.png": PNG,
            "word/_rels/vbaProject.bin.rels": RELS_EMPTY.replace(
                "</Relationships>",
                rel("rId1", f"{MS_REL}/wordVbaData", "vbaData.xml")
                + "</Relationships>",
            ),
        },
        edit={
            "word/document.xml": lambda xml: xml.replace(
                "<w:sectPr", WORD_OBJECTS + "<w:sectPr", 1
            ),
            "word/settings.xml": lambda xml: re.sub(
                r"(<w:settings[^>]*>)", r'\1<w:attachedTemplate r:id="rIdTpl1"/>', xml
            ).replace(
                "<w:settings ",
                f'<w:settings xmlns:r="{REL}" ',
                1 if "xmlns:r=" not in xml.split(">", 2)[1] else 0,
            ),
        },
        types=[
            override("word/vbaProject.bin", "application/vnd.ms-office.vbaProject"),
            override("word/vbaData.xml", "application/vnd.ms-word.vbaData+xml"),
            override(
                "word/activeX/activeX1.xml", "application/vnd.ms-office.activeX+xml"
            ),
            override("word/activeX/activeX1.bin", "application/vnd.ms-office.activeX"),
            override(
                "word/embeddings/oleObject1.bin",
                "application/vnd.openxmlformats-officedocument.oleObject",
            ),
            override("word/media/preview1.png", "image/png"),
        ],
        rels={
            "word/_rels/document.xml.rels": [
                rel("rIdVba1", f"{MS_REL}/vbaProject", "vbaProject.bin"),
                rel("rIdAx1", f"{REL}/control", "activeX/activeX1.xml"),
                rel("rIdOle1", f"{REL}/oleObject", "embeddings/oleObject1.bin"),
                rel("rIdImg1", f"{REL}/image", "media/preview1.png"),
                rel("rIdLink1", f"{REL}/hyperlink", "https://aelira.ai", True),
            ],
            "word/_rels/settings.xml.rels": [
                rel(
                    "rIdTpl1",
                    f"{REL}/attachedTemplate",
                    "https://example.com/template.dotm",
                    True,
                )
            ],
        },
        main_type="application/vnd.ms-word.document.macroEnabled.main+xml",
    )


REMOVED_WORD_PARTS = {
    "word/vbaProject.bin",
    "word/vbaData.xml",
    "word/_rels/vbaProject.bin.rels",
    "word/activeX/activeX1.xml",
    "word/activeX/activeX1.bin",
    "word/activeX/_rels/activeX1.xml.rels",
    "word/embeddings/oleObject1.bin",
}


@pytest.mark.asyncio
async def test_word_file_is_rebuilt_readable_without_active_content():
    original = macro_word_document()
    before = await DocumentValidator().validate("course.docx", original)
    assert before.threat_level == ThreatLevel.HIGH
    assert {"office_macros_detected", "office_external_content"} <= {
        f.category for f in before.findings
    }
    assert REMOVED_WORD_PARTS <= names(original)

    rebuilt, result = await sanitize_document("course.docx", original)

    assert result.sanitized is True
    assert result.threat_level.rank < ThreatLevel.HIGH.rank
    assert not names(rebuilt) & REMOVED_WORD_PARTS
    # Nothing still names a removed part.
    types_xml = read(rebuilt, "[Content_Types].xml")
    assert "vbaProject" not in types_xml and "activeX" not in types_xml
    assert "oleObject" not in types_xml
    assert "macroEnabled" not in types_xml
    targets = {attrs.get("Target") for _, attrs in all_relationships(rebuilt)}
    assert not targets & {
        "vbaProject.bin",
        "activeX/activeX1.xml",
        "embeddings/oleObject1.bin",
        "https://example.com/template.dotm",
    }
    document_xml = read(rebuilt, "word/document.xml")
    assert "w:control" not in document_xml and "OLEObject" not in document_xml
    assert "attachedTemplate" not in read(rebuilt, "word/settings.xml")
    # The equation preview, the hyperlink and the namespace prefixes stay.
    assert 'r:id="rIdImg1"' in document_xml
    assert ("rIdLink1", "External") in {
        (attrs.get("Id"), attrs.get("TargetMode"))
        for _, attrs in all_relationships(rebuilt)
    }
    original_root = read(original, "word/document.xml").split(">", 2)[1]
    assert read(rebuilt, "word/document.xml").split(">", 2)[1] == original_root

    opened = docx.Document(io.BytesIO(rebuilt))
    text = [paragraph.text for paragraph in opened.paragraphs]
    assert "Valve practical" in text
    assert "Record the reading before you open the valve." in text
    assert "Aelira website" in "".join(text)


@pytest.mark.asyncio
async def test_linked_ole_object_loading_a_web_page_is_removed():
    """A remote OLE link, as in CVE-2022-30190, with no macros at all."""
    original = edit_package(
        word_base(),
        edit={
            "word/document.xml": lambda xml: xml.replace(
                "<w:sectPr",
                '<w:p xmlns:o="urn:schemas-microsoft-com:office:office"><w:r>'
                '<w:object><o:OLEObject Type="Link" ProgID="htmlfile" '
                'ObjectID="_2" r:id="rIdRemote"/></w:object></w:r></w:p><w:sectPr',
                1,
            )
        },
        rels={
            "word/_rels/document.xml.rels": [
                rel(
                    "rIdRemote",
                    f"{REL}/oleObject",
                    "https://example.com/payload.html!",
                    True,
                )
            ]
        },
    )
    before = await DocumentValidator().validate("letter.docx", original)
    assert before.threat_level == ThreatLevel.HIGH
    assert "office_external_content" in {f.category for f in before.findings}

    rebuilt, result = await sanitize_document("letter.docx", original)

    assert result.sanitized is True
    assert "payload.html" not in str(all_relationships(rebuilt))
    assert "OLEObject" not in read(rebuilt, "word/document.xml")
    docx.Document(io.BytesIO(rebuilt))


@pytest.mark.asyncio
async def test_vba_project_under_another_name_is_removed():
    original = edit_package(
        word_base(),
        add={"word/payload.dat": vba_project_bin(BENIGN_VBA)},
        types=[override("word/payload.dat", "application/vnd.ms-office.vbaProject")],
        rels={
            "word/_rels/document.xml.rels": [
                rel("rIdVba1", f"{MS_REL}/vbaProject", "payload.dat")
            ]
        },
    )

    rebuilt, result = await sanitize_document("notes.docx", original)

    assert "word/payload.dat" in names(original)
    assert "word/payload.dat" not in names(rebuilt)
    assert result.sanitized is True
    docx.Document(io.BytesIO(rebuilt))


def workbook_base():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Data"
    workbook.active["A1"] = "Reading"
    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Total"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def macro_sheet_workbook():
    """Excel 4 macro sheet between two sheets, run by an Auto_Open name."""

    def add_sheet_and_names(xml):
        xml = re.sub(
            r'(<sheet [^>]*name="Data"[^>]*/>)',
            rf'\1<sheet xmlns:r="{REL}" name="Macro1" sheetId="9" r:id="rIdMacro1"/>',
            xml,
            count=1,
        )
        names_xml = (
            "<definedNames>"
            '<definedName name="_xlnm.Auto_Open">Macro1!$A$1</definedName>'
            '<definedName name="_xlnm.Print_Area" localSheetId="2">'
            "Summary!$A$1:$B$2</definedName>"
            "</definedNames>"
        )
        xml = re.sub(r"<definedNames>.*?</definedNames>|<definedNames/>", "", xml)
        return xml.replace("</sheets>", "</sheets>" + names_xml, 1)

    return edit_package(
        workbook_base(),
        add={
            "xl/macrosheets/sheet1.xml": '<xm:macrosheet xmlns:xm="http://schemas.'
            'microsoft.com/office/excel/2006/main"/>',
        },
        edit={"xl/workbook.xml": add_sheet_and_names},
        types=[
            override(
                "xl/macrosheets/sheet1.xml", "application/vnd.ms-excel.macrosheet+xml"
            )
        ],
        rels={
            "xl/_rels/workbook.xml.rels": [
                rel("rIdMacro1", f"{MS_REL}/xlMacrosheet", "macrosheets/sheet1.xml")
            ]
        },
    )


@pytest.mark.asyncio
async def test_excel_4_macro_sheet_is_detected_and_removed():
    original = macro_sheet_workbook()
    before = await DocumentValidator().validate("marks.xlsx", original)
    assert before.threat_level == ThreatLevel.HIGH
    assert "office_xlm_macros" in {f.category for f in before.findings}

    rebuilt, result = await sanitize_document("marks.xlsx", original)

    assert result.sanitized is True
    assert "xl/macrosheets/sheet1.xml" not in names(rebuilt)
    workbook_xml = read(rebuilt, "xl/workbook.xml")
    assert "Macro1" not in workbook_xml and "Auto_Open" not in workbook_xml
    # The print area still belongs to Summary, now the second sheet.
    assert 'localSheetId="1"' in workbook_xml
    workbook = openpyxl.load_workbook(io.BytesIO(rebuilt))
    assert workbook.sheetnames == ["Data", "Summary"]
    assert workbook["Data"]["A1"].value == "Reading"
    assert workbook["Summary"].print_area is not None


@pytest.mark.asyncio
async def test_workbook_with_vba_is_rebuilt_as_xlsx():
    original = edit_package(
        workbook_base(),
        add={"xl/vbaProject.bin": vba_project_bin(BENIGN_VBA)},
        types=[override("xl/vbaProject.bin", "application/vnd.ms-office.vbaProject")],
        rels={
            "xl/_rels/workbook.xml.rels": [
                rel("rIdVba1", f"{MS_REL}/vbaProject", "vbaProject.bin")
            ]
        },
        main_type="application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    )

    rebuilt, result = await sanitize_document("marks.xlsx", original)

    assert result.sanitized is True
    assert "xl/vbaProject.bin" not in names(rebuilt)
    assert "macroEnabled" not in read(rebuilt, "[Content_Types].xml")
    assert openpyxl.load_workbook(io.BytesIO(rebuilt)).sheetnames == ["Data", "Summary"]


@pytest.mark.asyncio
async def test_presentation_with_vba_is_rebuilt_as_pptx():
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Week 3 lab"
    buffer = io.BytesIO()
    presentation.save(buffer)
    original = edit_package(
        buffer.getvalue(),
        add={"ppt/vbaProject.bin": vba_project_bin(BENIGN_VBA)},
        types=[override("ppt/vbaProject.bin", "application/vnd.ms-office.vbaProject")],
        rels={
            "ppt/_rels/presentation.xml.rels": [
                rel("rIdVba1", f"{MS_REL}/vbaProject", "vbaProject.bin")
            ]
        },
        main_type="application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml",
    )

    rebuilt, result = await sanitize_document("lab.pptx", original)

    assert result.sanitized is True
    assert "ppt/vbaProject.bin" not in names(rebuilt)
    opened = pptx.Presentation(io.BytesIO(rebuilt))
    assert opened.slides[0].shapes.title.text == "Week 3 lab"


@pytest.mark.parametrize(
    "broken",
    ["doctype", "entity", "unreadable_rels"],
)
def test_rebuild_refuses_package_xml_it_cannot_trust(broken):
    original = word_base()
    if broken == "doctype":
        change = {
            "word/_rels/document.xml.rels": lambda xml: xml.replace(
                "<Relationships", '<!DOCTYPE r [<!ENTITY x "y">]><Relationships', 1
            )
        }
    elif broken == "entity":
        change = {
            "word/_rels/document.xml.rels": lambda xml: xml.replace(
                "<Relationships",
                '<!ENTITY x SYSTEM "file:///etc/passwd"><Relationships',
                1,
            )
        }
    else:
        change = {"word/_rels/document.xml.rels": lambda xml: xml[:-20]}

    with pytest.raises(SanitizationError):
        rebuild_ooxml(edit_package(original, edit=change))


@pytest.mark.asyncio
async def test_legacy_office_file_is_refused_not_returned_unchanged(monkeypatch):
    monkeypatch.setattr(
        "src.config.settings.get_settings",
        lambda: type("S", (), {"env": "test", "max_file_size_pdf": 10**7})(),
    )
    legacy = vba_project_bin(BENIGN_VBA)
    assert (await DocumentValidator().validate("a.doc", legacy)).threat_level == (
        ThreatLevel.HIGH
    )

    with pytest.raises(SanitizationError):
        await sanitize_document("a.doc", legacy)


class MemorySession:
    def __init__(self):
        self.added = []

    def add(self, value):
        self.added.append(value)

    def commit(self):
        pass


def upload(content, filename):
    from fastapi import UploadFile

    return UploadFile(file=io.BytesIO(content), filename=filename)


@pytest.fixture
def session(monkeypatch):
    monkeypatch.setattr(
        "src.config.settings.get_settings",
        lambda: type("S", (), {"env": "test", "max_file_size_pdf": 10**7})(),
    )
    return MemorySession()


@pytest.mark.asyncio
async def test_upload_of_word_file_with_macros_returns_the_rebuilt_copy(session):
    original = macro_word_document()

    returned = await _shared.validate_uploaded_file(
        upload(original, "course.docx"), session, "workspace"
    )

    assert "word/vbaProject.bin" not in names(returned)
    docx.Document(io.BytesIO(returned))
    record = session.added[0]
    assert record.was_sanitized is True
    assert record.was_blocked is False


@pytest.mark.asyncio
async def test_upload_whose_rebuild_fails_is_refused_and_recorded(session, monkeypatch):
    def fail(content):
        raise SanitizationError("package_xml_unreadable")

    monkeypatch.setattr("src.security.document_validator.rebuild_ooxml", fail)

    with pytest.raises(HTTPException) as caught:
        await _shared.validate_uploaded_file(
            upload(macro_word_document(), "course.docx"), session, "workspace"
        )

    assert caught.value.status_code == 400
    assert caught.value.detail["threat_level"] == "high"
    assert "could not be rebuilt" in caught.value.detail["message"]
    record = session.added[0]
    assert record.was_blocked is True
    assert record.was_sanitized is False
    assert record.blocked_reason == "Document contains VBA macros"


@pytest.mark.asyncio
async def test_high_code_file_is_processed_unchanged_and_not_recorded_as_sanitised(
    session,
):
    page = b"<html><body><script>document.cookie = 'a=1';</script></body></html>"
    result = await DocumentValidator().validate("page.html", page)
    assert result.threat_level == ThreatLevel.HIGH

    returned = await _shared.validate_uploaded_file(
        upload(page, "page.html"), session, "workspace"
    )

    assert returned == page
    assert session.added[0].was_sanitized is False
    assert session.added[0].was_blocked is False


def scripted_pdf(open_action=None):
    pdf = pikepdf.new()
    pdf.add_blank_page(page_size=(300, 300))
    page = pdf.pages[0]
    page.obj.Resources = pikepdf.Dictionary(
        Font=pikepdf.Dictionary(
            F1=pikepdf.Dictionary(
                Type=pikepdf.Name.Font,
                Subtype=pikepdf.Name.Type1,
                BaseFont=pikepdf.Name.Helvetica,
            )
        )
    )
    page.obj.Contents = pdf.make_stream(b"BT /F1 12 Tf 20 250 Td (Unit outline) Tj ET")

    def link(action):
        return pdf.make_indirect(
            pikepdf.Dictionary(
                Type=pikepdf.Name.Annot,
                Subtype=pikepdf.Name.Link,
                Rect=[20, 20, 120, 40],
                A=action,
            )
        )

    page.obj.Annots = pdf.make_indirect(
        pikepdf.Array(
            [
                link(pikepdf.Dictionary(S=pikepdf.Name.URI, URI="https://aelira.ai")),
                link(
                    pikepdf.Dictionary(
                        S=pikepdf.Name.JavaScript, JS=pikepdf.String("app.alert(1)")
                    )
                ),
            ]
        )
    )
    page.obj.AA = pikepdf.Dictionary(
        O=pikepdf.Dictionary(S=pikepdf.Name.JavaScript, JS=pikepdf.String("x = 1"))
    )
    pdf.Root.Names = pikepdf.Dictionary(
        JavaScript=pikepdf.Dictionary(
            Names=pikepdf.Array(
                [
                    pikepdf.String("init"),
                    pikepdf.Dictionary(
                        S=pikepdf.Name.JavaScript, JS=pikepdf.String("y = 2")
                    ),
                ]
            )
        )
    )
    pdf.Root.OpenAction = open_action or pikepdf.Dictionary(
        S=pikepdf.Name.JavaScript, JS=pikepdf.String("this.zoom = 100")
    )
    buffer = io.BytesIO()
    pdf.save(buffer, object_stream_mode=pikepdf.ObjectStreamMode.disable)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_pdf_is_rebuilt_without_scripts_and_keeps_links_and_text():
    original = scripted_pdf()
    before = await DocumentValidator().validate("outline.pdf", original)
    assert before.threat_level == ThreatLevel.HIGH

    rebuilt, result = await sanitize_document("outline.pdf", original)

    assert result.sanitized is True
    assert result.threat_level.rank < ThreatLevel.HIGH.rank
    for token in (b"/JavaScript", b"/JS", b"/AA", b"/OpenAction"):
        assert token not in rebuilt
    with pikepdf.open(io.BytesIO(rebuilt)) as pdf:
        actions = [str(a.A.S) for a in pdf.pages[0].Annots if "/A" in a]
        assert actions == ["/URI"]
    with fitz.open(stream=rebuilt, filetype="pdf") as document:
        assert "Unit outline" in document[0].get_text()


def test_pdf_open_destination_is_kept():
    with pikepdf.open(io.BytesIO(scripted_pdf())) as source:
        destination = pikepdf.Array([source.pages[0].obj, pikepdf.Name.Fit])
        source.Root.OpenAction = destination
        buffer = io.BytesIO()
        source.save(buffer)

    rebuilt = rebuild_pdf(buffer.getvalue())

    with pikepdf.open(io.BytesIO(rebuilt)) as pdf:
        assert isinstance(pdf.Root.OpenAction, pikepdf.Array)


def test_unreadable_pdf_is_refused():
    with pytest.raises(SanitizationError):
        rebuild_pdf(b"%PDF-1.7\nnot really a pdf")


@pytest.mark.asyncio
async def test_pdf_names_are_matched_whole_and_with_escapes_decoded():
    validator = DocumentValidator()
    escaped = b"%PDF-1.4\n1 0 obj << /S /J#61vaScript /J#53 (x) >> endobj\n%%EOF"
    lookalike = b"%PDF-1.4\n1 0 obj << /JSomething 1 /AAPL 2 >> endobj\n%%EOF"

    assert (await validator.validate("a.pdf", escaped)).threat_level == (
        ThreatLevel.HIGH
    )
    found = {f.details for f in (await validator.validate("a.pdf", lookalike)).findings}
    assert not any("/JS" in (d or "") or "/AA" in (d or "") for d in found)


@pytest.mark.asyncio
async def test_external_subdocument_is_detected_and_removed():
    original = edit_package(
        word_base(),
        edit={
            "word/document.xml": lambda xml: xml.replace(
                "<w:sectPr", '<w:subDoc r:id="rIdSub1"/><w:sectPr', 1
            )
        },
        rels={
            "word/_rels/document.xml.rels": [
                rel(
                    "rIdSub1",
                    f"{REL}/subDocument",
                    "https://example.com/part.docx",
                    True,
                )
            ]
        },
    )
    before = await DocumentValidator().validate("thesis.docx", original)
    assert "office_external_content" in {f.category for f in before.findings}

    rebuilt, result = await sanitize_document("thesis.docx", original)

    assert result.sanitized is True
    assert "subDoc" not in read(rebuilt, "word/document.xml")
    assert "part.docx" not in str(all_relationships(rebuilt))
    docx.Document(io.BytesIO(rebuilt))


@pytest.mark.asyncio
async def test_presentation_ole_object_frame_is_removed_with_its_object():
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Results"
    slide.shapes.add_ole_object(
        io.BytesIO(workbook_base()),
        "Excel.Sheet.12",
        Inches(1),
        Inches(2),
        icon_file=io.BytesIO(PNG),
    )
    buffer = io.BytesIO()
    presentation.save(buffer)
    original = edit_package(
        buffer.getvalue(),
        add={"ppt/vbaProject.bin": vba_project_bin(BENIGN_VBA)},
        types=[override("ppt/vbaProject.bin", "application/vnd.ms-office.vbaProject")],
        rels={
            "ppt/_rels/presentation.xml.rels": [
                rel("rIdVba1", f"{MS_REL}/vbaProject", "vbaProject.bin")
            ]
        },
    )
    assert any(name.startswith("ppt/embeddings/") for name in names(original))

    rebuilt, result = await sanitize_document("results.pptx", original)

    assert result.sanitized is True
    assert not any(name.startswith("ppt/embeddings/") for name in names(rebuilt))
    assert "oleObj" not in read(rebuilt, "ppt/slides/slide1.xml")
    shapes = pptx.Presentation(io.BytesIO(rebuilt)).slides[0].shapes
    assert [shape.has_text_frame and shape.text_frame.text for shape in shapes] == [
        "Results"
    ]


def test_detection_reads_only_package_xml_and_rebuild_is_size_bounded(monkeypatch):
    from src.security import active_content

    original = macro_word_document()
    monkeypatch.setattr(active_content, "_MAX_PART_BYTES", 4096)

    # Detection reads only content types and relationships, which are small.
    assert active_content.find_active_content(original) is not None
    # A rebuild reads every part, and refuses a part over the limit.
    with pytest.raises(SanitizationError) as caught:
        rebuild_ooxml(original)
    assert caught.value.reason == "package_too_large"


@pytest.mark.asyncio
async def test_pdf_name_followed_by_nul_is_found():
    """NUL is PDF whitespace but not a Python regex \\s."""
    content = b"%PDF-1.4\n1 0 obj << /S /JavaScript\x00>> endobj\n%%EOF"

    result = await DocumentValidator().validate("a.pdf", content)

    assert result.threat_level == ThreatLevel.HIGH


@pytest.mark.asyncio
async def test_pdf_script_inside_a_compressed_object_stream_is_found():
    with pikepdf.open(io.BytesIO(scripted_pdf())) as pdf:
        buffer = io.BytesIO()
        pdf.save(buffer, object_stream_mode=pikepdf.ObjectStreamMode.generate)
    hidden = buffer.getvalue()
    assert b"/JavaScript" not in hidden and b"/JS" not in hidden

    result = await DocumentValidator().validate("outline.pdf", hidden)
    assert result.threat_level == ThreatLevel.HIGH

    rebuilt, rebuilt_result = await sanitize_document("outline.pdf", hidden)
    assert rebuilt_result.sanitized is True
    assert rebuilt_result.threat_level.rank < ThreatLevel.HIGH.rank


@pytest.mark.asyncio
async def test_only_code_pattern_findings_pass_through_without_a_rebuild():
    page = b"<html><body><script>document.cookie = 'a=1';</script></body></html>"
    passed, result = await sanitize_document("page.html", page)
    assert passed == page and result.sanitized is False

    bundle = io.BytesIO()
    with zipfile.ZipFile(bundle, "w") as zf:
        zf.writestr("site/index.html", page.decode())
    passed, result = await sanitize_document("site.zip", bundle.getvalue())
    assert result.sanitized is False

    # A zip uploaded as .html is a type mismatch, not a code pattern.
    with pytest.raises(SanitizationError):
        await sanitize_document("page.html", bundle.getvalue())


@pytest.mark.asyncio
async def test_xhtml_is_not_an_extension_mismatch():
    page = b'<?xml version="1.0"?><html xmlns="http://www.w3.org/1999/xhtml"></html>'

    result = await DocumentValidator().validate("page.html", page)

    assert "file_type_mismatch" not in {f.category for f in result.findings}


def clean_pdf(object_streams):
    with pikepdf.open(io.BytesIO(scripted_pdf())) as pdf:
        pdf.Root.OpenAction = pikepdf.Array([pdf.pages[0].obj, pikepdf.Name.Fit])
        del pdf.Root["/Names"]
        del pdf.pages[0].obj["/AA"]
        pdf.pages[0].obj.Annots = pikepdf.Array([])
        buffer = io.BytesIO()
        mode = (
            pikepdf.ObjectStreamMode.generate
            if object_streams
            else pikepdf.ObjectStreamMode.disable
        )
        pdf.save(buffer, object_stream_mode=mode)
    return buffer.getvalue()


@pytest.mark.asyncio
@pytest.mark.parametrize("object_streams", [True, False])
async def test_incomplete_pdf_walk_fails_closed(monkeypatch, object_streams):
    from src.security import document_validator

    content = clean_pdf(object_streams)
    validator = DocumentValidator()
    assert (await validator.validate("a.pdf", content)).threat_level.rank < (
        ThreatLevel.HIGH.rank
    )

    monkeypatch.setattr(document_validator, "_MAX_PDF_OBJECTS", 1)
    result = await validator.validate("a.pdf", content)
    assert "pdf_structure_unchecked" in {f.category for f in result.findings}
    assert result.threat_level == ThreatLevel.HIGH

    # A rebuild that cannot be walked either is refused.
    with pytest.raises(SanitizationError):
        await sanitize_document("a.pdf", content)


@pytest.mark.asyncio
async def test_pdf_that_cannot_be_parsed_is_high():
    result = await DocumentValidator().validate("a.pdf", b"%PDF-1.4\nnot a pdf")

    assert "pdf_structure_unchecked" in {f.category for f in result.findings}
    assert result.threat_level == ThreatLevel.HIGH
