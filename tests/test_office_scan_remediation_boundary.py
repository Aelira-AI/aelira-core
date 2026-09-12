"""Real Office scan persistence must remain usable at the worker boundary."""

from copy import deepcopy
from types import SimpleNamespace
from unittest.mock import MagicMock

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from test_office_measured_remediation import office_fixture
from src.jobs.remediation_subprocess import _build_remediator


def persist_scan(tmp_path, monkeypatch, kind):
    from src.api.education import scan_routes
    from src.db import database
    from src.ai import workspace_provider_runtime
    from src.middleware import quota

    source, scan, _ = office_fixture(tmp_path, kind)
    if kind == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = "Course overview"
        slide.shapes.title.fill.solid()
        slide.shapes.title.fill.fore_color.rgb = RGBColor(255, 255, 255)
        slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(
            150, 150, 150
        )
        presentation.save(source)
    scan_row = SimpleNamespace(id="scan-test", department_id="workspace-test")
    db = MagicMock()
    db.query.return_value.filter.return_value.first.return_value = scan_row
    monkeypatch.setattr(database, "SessionLocal", lambda: db)
    monkeypatch.setattr(
        workspace_provider_runtime, "workspace_provider_runtime", lambda _: None
    )
    monkeypatch.setattr(quota, "increment_usage_sync", lambda *a, **k: None)
    uploaded = tmp_path / f"upload.{kind}"
    uploaded.write_bytes(source.read_bytes())
    getattr(scan_routes, f"process_{kind}_background")(
        str(uploaded),
        source.read_bytes(),
        source.name,
        scan_row.id,
        False,
        False,
        str(source),
        "user-test",
        "workspace-test",
    )
    assert db.add.call_count == 1, scan_row.__dict__
    stored = db.add.call_args.args[0]
    return source, scan, stored


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_persisted_findings_produce_measured_saved_output(tmp_path, monkeypatch, kind):
    source, scan, stored = persist_scan(tmp_path, monkeypatch, kind)
    before = scan(str(source))
    assert stored.issues == before.issues
    assert sum(
        getattr(stored, f"{severity}_issues")
        for severity in ("critical", "high", "medium", "low")
    ) == len(before.issues)
    original = deepcopy(stored.issues)
    worker = _build_remediator(
        {"scan_type": kind, "issues": stored.issues, "options": {"use_ai": False}},
        source,
        tmp_path / "output",
    )
    result = worker.remediate()
    assert result.success, result.error_message
    assert result.fixed_count > 0, (before.issues, result.model_dump())
    after = scan(result.output_file)
    assert result.original_compliance_score == before.compliance_score
    assert result.remediated_compliance_score == after.compliance_score
    assert result.remediated_compliance_score > result.original_compliance_score
    assert result.score_provenance == "scanner_rescan"
    assert (result.original_compliance_score, result.remediated_compliance_score) == {
        "docx": (88, 92),
        "pptx": (92, 100),
        "xlsx": (49, 100),
    }[kind]
    assert stored.issues == original


def legacy_rows(canonical, kind):
    rows = []
    for index, issue in enumerate(canonical):
        raw = deepcopy(issue["metadata"])
        raw["type"] = raw.pop("scanner_type")
        if "scanner_issue_type" in raw:
            raw["issue_type"] = raw.pop("scanner_issue_type")
        if kind == "pptx":
            raw["slide"] = raw.pop("slide_index") + 1
        raw.update(id=f"retained-{index}", severity="medium")
        rows.append(raw)
    return rows


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_legacy_rows_recovered_from_real_original(tmp_path, monkeypatch, kind):
    source, scan, stored = persist_scan(tmp_path, monkeypatch, kind)
    issues = legacy_rows(stored.issues, kind)
    original = deepcopy(issues)
    result = _build_remediator(
        {"scan_type": kind, "issues": issues, "options": {"use_ai": False}},
        source,
        tmp_path / "legacy-output",
    ).remediate()
    assert result.success
    assert result.fixed_count > 0
    assert result.original_compliance_score == scan(str(source)).compliance_score
    assert (
        result.remediated_compliance_score == scan(result.output_file).compliance_score
    )
    assert result.improvement > 0
    assert {fix.issue_id for fix in result.fixed_issues} <= {
        row["id"] for row in issues
    }
    assert issues == original


@pytest.mark.parametrize("approved", [False, True])
def test_legacy_subset_preserves_identity_and_review_content(
    tmp_path, monkeypatch, approved
):
    source, scan, stored = persist_scan(tmp_path, monkeypatch, "docx")
    heading = legacy_rows(stored.issues, "docx")[0]
    heading["fixed_content"] = "1"
    worker = _build_remediator(
        {
            "scan_type": "docx",
            "issues": [heading],
            "options": {"use_ai": False, "approved_fixes_only": approved},
        },
        source,
        tmp_path / "subset-output",
    )
    assert len(worker.issues) == 1
    assert worker.issues[0].id == heading["id"]
    assert worker.issues[0].metadata["fixed_content"] == "1"
    result = worker.remediate()
    assert result.original_compliance_score == 88
    assert result.fixed_count == (0 if approved else 1)
    assert result.remediated_compliance_score == (88 if approved else 89)
    assert (
        result.remediated_compliance_score == scan(result.output_file).compliance_score
    )


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_anchor",
        "stale_anchor",
        "duplicate",
        "purpose_conflict",
        "unknown_category",
        "bad_metadata",
        "semantic",
    ],
)
def test_unreproducible_legacy_findings_stay_unavailable(
    tmp_path, monkeypatch, mutation
):
    source, _, stored = persist_scan(tmp_path, monkeypatch, "docx")
    rows = legacy_rows(stored.issues, "docx")[:1]
    if mutation == "missing_anchor":
        rows[0] = {"type": "heading", "issue_type": "missing_h1"}
    elif mutation == "stale_anchor":
        rows[0]["paragraph_index"] = 900
    elif mutation == "duplicate":
        rows.append(deepcopy(rows[0]))
    elif mutation == "purpose_conflict":
        rows[0]["rule_id"] = "alt_text"
    elif mutation == "unknown_category":
        rows[0]["metadata"] = {"category": "unrecognized"}
    elif mutation == "bad_metadata":
        rows[0]["metadata"] = "malformed"
    else:
        rows[0] = {
            "type": "image",
            "image_index": 0,
            "has_alt_text": True,
            "alt_text_validated": True,
            "description": "Inaccurate alt text requires review",
        }
    result = _build_remediator(
        {"scan_type": "docx", "issues": rows, "options": {"use_ai": False}},
        source,
        tmp_path / "refusal-output",
    ).remediate()
    assert result.fixed_count == 0
    assert result.remediated_compliance_score is None
    assert result.manual_count > 0


def test_missing_original_cannot_recover_legacy_findings(tmp_path):
    with pytest.raises(Exception):
        _build_remediator(
            {
                "scan_type": "docx",
                "issues": [{"type": "heading", "paragraph_index": 0}],
                "options": {"use_ai": False},
            },
            tmp_path / "missing.docx",
            tmp_path / "output",
        )


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_supplied_worker_findings_preserve_exact_content_and_subset(
    tmp_path, monkeypatch, kind
):
    from docx import Document
    from openpyxl import load_workbook

    source, _, stored = persist_scan(tmp_path, monkeypatch, kind)
    selected = (
        next(issue for issue in stored.issues if issue["category"] == "title")
        if kind == "docx"
        else stored.issues[0]
    )
    content = {
        "docx": "Approved course guide",
        "pptx": "#000000",
        "xlsx": "Approved courses",
    }[kind]
    # Exercise the worker issue representation constructed by remediation_job.
    # Public enqueue routes currently authorize reviewed-fix jobs for CODE only;
    # this checks Office worker safety without extending that route policy.
    issues = [
        {
            "id": "reviewed-issue",
            "category": selected["category"],
            "severity": selected["severity"],
            "description": selected["description"],
            "location": selected["location"],
            "original_content": None,
            "fix_suggestion": content,
            "fixed_content": content,
            "wcag_criteria": selected.get("wcag_criteria"),
            "metadata": {},
        }
    ]
    original = deepcopy(issues)
    result = _build_remediator(
        {"scan_type": kind, "issues": issues, "options": {"approved_fixes_only": True}},
        source,
        tmp_path / "review-output",
    ).remediate()
    assert issues == original
    assert result.total_issues == 1
    assert result.fixed_count == 1, result
    assert result.fixed_issues[0].issue_id == "reviewed-issue"
    assert result.fixed_issues[0].fixed_content == content
    if kind == "docx":
        assert Document(result.output_file).core_properties.title == content
    elif kind == "pptx":
        assert (
            str(
                Presentation(result.output_file)
                .slides[0]
                .shapes.title.text_frame.paragraphs[0]
                .runs[0]
                .font.color.rgb
            )
            == "000000"
        )
    else:
        assert load_workbook(result.output_file).sheetnames == [content]


def test_nested_semantic_evidence_survives_legacy_recovery(tmp_path, monkeypatch):
    source, _, stored = persist_scan(tmp_path, monkeypatch, "docx")
    from io import BytesIO
    from PIL import Image
    from docx import Document
    from src.education.docx_processor import DocxProcessor

    image = BytesIO()
    Image.new("RGB", (20, 20), "red").save(image, format="PNG")
    image.seek(0)
    document = Document(source)
    document.add_picture(image)
    document.save(source)
    # A nested semantic claim must survive even when deterministic metadata has
    # a different value; it may never be downgraded into a presence-only pass.
    rows = [
        row
        for row in legacy_rows(DocxProcessor().process_docx(str(source)).issues, "docx")
        if row["type"] == "image"
    ]
    rows[0]["metadata"] = {"alt_text_validated": True, "alt_text_accurate": False}
    worker = _build_remediator(
        {"scan_type": "docx", "issues": rows, "options": {"use_ai": False}},
        source,
        tmp_path / "output",
    )
    assert worker.issues[0].metadata["alt_text_validated"] is True
    assert worker.issues[0].metadata["alt_text_accurate"] is False
    result = worker.remediate()
    assert result.remediated_compliance_score is None
    assert result.fixed_count == 0


def test_empty_slide_does_not_get_fabricated_numbered_title(tmp_path):
    source, scan, _ = office_fixture(tmp_path, "pptx")
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(source)
    result = _build_remediator(
        {
            "scan_type": "pptx",
            "issues": scan(str(source)).issues,
            "options": {"use_ai": False},
        },
        source,
        tmp_path / "output",
    ).remediate()
    assert result.fixed_count == 0
    assert result.manual_count == 1
    assert not any(
        shape.has_text_frame and shape.text == "Slide 1"
        for shape in Presentation(result.output_file).slides[0].shapes
    )


def test_reviewed_heading_cannot_be_replaced_with_different_rule_action(
    tmp_path, monkeypatch
):
    from docx import Document

    source, _, stored = persist_scan(tmp_path, monkeypatch, "docx")
    heading = deepcopy(stored.issues[0])
    heading.update(fixed_content="Heading 3", fix_suggestion="Heading 3")
    result = _build_remediator(
        {
            "scan_type": "docx",
            "issues": [heading],
            "options": {"approved_fixes_only": True},
        },
        source,
        tmp_path / "output",
    ).remediate()
    assert result.fixed_count == 0
    assert result.manual_count == 1
    assert Document(result.output_file).paragraphs[0].style.name == "Heading 2"
    assert result.remediated_compliance_score == 88


def test_reviewed_sheet_name_collision_is_not_silently_changed(tmp_path):
    from openpyxl import load_workbook
    from src.education.xlsx_processor import XlsxProcessor

    source, _, _ = office_fixture(tmp_path, "xlsx")
    workbook = load_workbook(source)
    workbook.create_sheet("Course")
    workbook.save(source)
    issue = next(
        issue
        for issue in XlsxProcessor().process_xlsx(str(source)).issues
        if issue["category"] == "sheet_name"
    )
    issue.update(fixed_content="Course", fix_suggestion="Course")
    result = _build_remediator(
        {
            "scan_type": "xlsx",
            "issues": [issue],
            "options": {"approved_fixes_only": True},
        },
        source,
        tmp_path / "output",
    ).remediate()
    assert result.fixed_count == 0
    assert load_workbook(result.output_file).sheetnames == ["Sheet", "Course"]


def test_partial_contrast_improvement_does_not_claim_fix(tmp_path, monkeypatch):
    source, scan, _ = persist_scan(tmp_path, monkeypatch, "pptx")
    presentation = Presentation(source)
    presentation.slides[0].shapes.title.text_frame.paragraphs[0].runs[
        0
    ].font.color.rgb = RGBColor(180, 180, 180)
    presentation.save(source)
    result = _build_remediator(
        {
            "scan_type": "pptx",
            "issues": scan(str(source)).issues,
            "options": {"use_ai": False},
        },
        source,
        tmp_path / "output",
    ).remediate()
    assert result.fixed_count == 0
    assert result.manual_count == 1
    assert result.original_compliance_score == 92
    assert result.remediated_compliance_score == 97
    assert (
        result.remediated_compliance_score == scan(result.output_file).compliance_score
    )
