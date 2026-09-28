"""The bounded PowerPoint repair must change only the accepted saved order."""

import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches

from src.education.remediation.base import RemediationConfig
from src.education.remediation.pptx_reading_order import (
    UnsupportedReadingOrder,
    repair_saved_order,
    source_sha256,
    validate_order,
    verify_saved_order,
)
from src.education.remediation.pptx_remediator import PptxRemediator


def _deck(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(2), Inches(0.5))
    first.text = "Second in intended reading order"
    first.text_frame.paragraphs[0].runs[
        0
    ].hyperlink.address = "https://example.org/course"
    second = slide.shapes.add_textbox(Inches(4), Inches(0.5), Inches(2), Inches(0.5))
    second.text = "First in intended reading order"
    image = Image.new("RGB", (4, 4), (40, 80, 120))
    data = BytesIO()
    image.save(data, format="PNG")
    data.seek(0)
    picture = slide.shapes.add_picture(
        data, Inches(0.5), Inches(2), Inches(1), Inches(1)
    )
    slide.notes_slide.notes_text_frame.text = "Instructor notes remain intact."
    other_slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    other_slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text = (
        "Other slide"
    )
    presentation.save(path)
    return [first.shape_id, second.shape_id, picture.shape_id]


def _parts(path):
    with ZipFile(path) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _issue(source, target, slide_index=0):
    return {
        "id": f"reading-order-{slide_index}",
        "category": "reading_order",
        "severity": "medium",
        "description": "Reviewer accepted a slide object order",
        "location": f"Slide {slide_index + 1}",
        "metadata": {
            "slide_index": slide_index,
            "accepted_shape_ids": target,
            "source_sha256": source_sha256(source),
        },
    }


def test_saved_order_and_unrelated_package_parts_are_preserved(tmp_path):
    source, output = tmp_path / "input.pptx", tmp_path / "output.pptx"
    ids = _deck(source)
    original_bytes = source.read_bytes()
    target = [ids[1], ids[0], ids[2]]
    repair_saved_order(source, output, source_sha256(source), {0: target})

    reopened = Presentation(output)
    assert [shape.shape_id for shape in reopened.slides[0].shapes] == target
    assert reopened.slides[0].shapes[0].text == "First in intended reading order"
    assert (
        reopened.slides[0].shapes[1].text_frame.paragraphs[0].runs[0].hyperlink.address
        == "https://example.org/course"
    )
    assert (
        reopened.slides[0].notes_slide.notes_text_frame.text
        == "Instructor notes remain intact."
    )
    assert source.read_bytes() == original_bytes
    original, saved = _parts(source), _parts(output)
    assert original.keys() == saved.keys()
    assert [name for name in original if original[name] != saved[name]] == [
        "ppt/slides/slide1.xml"
    ]
    assert (
        hashlib.sha256(saved["ppt/media/image1.png"]).digest()
        == hashlib.sha256(original["ppt/media/image1.png"]).digest()
    )


def test_direct_reviewer_order_is_verified_after_delivery(tmp_path):
    source = tmp_path / "input.pptx"
    ids = _deck(source)
    result = PptxRemediator(
        str(source),
        [_issue(source, [ids[1], ids[0], ids[2]])],
        RemediationConfig(use_ai=False, use_supplied_fixes=True, create_backup=False),
    ).remediate()
    assert result.success, result.error_message
    assert result.fixed_count == 1
    assert result.fixed_issues[0].verification_passed
    assert result.fixed_issues[0].needs_review
    assert result.verification_passed
    assert [
        shape.shape_id for shape in Presentation(result.output_file).slides[0].shapes
    ] == [ids[1], ids[0], ids[2]]


def test_unreviewed_and_stale_orders_remain_manual(tmp_path):
    source = tmp_path / "input.pptx"
    ids = _deck(source)
    target = [ids[1], ids[0], ids[2]]
    for config, issue in [
        (
            RemediationConfig(use_ai=False, create_backup=False),
            {**_issue(source, target), "fixed_content": str(target)},
        ),
        (
            RemediationConfig(
                use_ai=False, use_supplied_fixes=True, create_backup=False
            ),
            {
                **_issue(source, target),
                "metadata": {
                    **_issue(source, target)["metadata"],
                    "source_sha256": "0" * 64,
                },
            },
        ),
    ]:
        result = PptxRemediator(str(source), [issue], config).remediate()
        assert result.fixed_count == 0
        assert result.manual_count == 1
        assert Path(result.output_file).read_bytes() == source.read_bytes()
    assert (
        "[Accessibility Note"
        not in Presentation(source).slides[0].notes_slide.notes_text_frame.text
    )


def test_overlapping_and_grouped_shapes_are_refused_without_output(tmp_path):
    source, output = tmp_path / "input.pptx", tmp_path / "output.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    first = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(2)
    )
    second = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(2), Inches(2), Inches(2), Inches(2)
    )
    presentation.save(source)
    with pytest.raises(UnsupportedReadingOrder, match="overlapping"):
        repair_saved_order(
            source,
            output,
            source_sha256(source),
            {0: [second.shape_id, first.shape_id]},
        )
    assert not output.exists()

    slide.shapes.add_group_shape([first, second])
    presentation.save(source)
    with pytest.raises(UnsupportedReadingOrder, match="Grouped"):
        validate_order(
            source,
            source_sha256(source),
            {0: [shape.shape_id for shape in slide.shapes]},
        )


def test_multiple_slides_and_already_correct_order(tmp_path):
    source, output = tmp_path / "input.pptx", tmp_path / "output.pptx"
    ids = _deck(source)
    presentation = Presentation(source)
    slide = presentation.slides[1]
    later = slide.shapes.add_textbox(Inches(5), Inches(2), Inches(1), Inches(1))
    slide_ids = [shape.shape_id for shape in slide.shapes]
    presentation.save(source)
    targets = {0: [ids[1], ids[0], ids[2]], 1: [later.shape_id, slide_ids[0]]}
    repair_saved_order(source, output, source_sha256(source), targets)
    reopened = Presentation(output)
    assert [shape.shape_id for shape in reopened.slides[0].shapes] == targets[0]
    assert [shape.shape_id for shape in reopened.slides[1].shapes] == targets[1]

    result = PptxRemediator(
        str(output),
        [_issue(output, targets[0])],
        RemediationConfig(use_ai=False, use_supplied_fixes=True, create_backup=False),
    ).remediate()
    assert result.fixed_count == 0
    assert result.manual_count == 1


def test_title_placeholders_and_animation_boundary(tmp_path):
    source, output = tmp_path / "input.pptx", tmp_path / "output.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = "Title"
    slide.placeholders[1].text = "Subtitle"
    presentation.save(source)
    ids = [shape.shape_id for shape in slide.shapes]
    repair_saved_order(source, output, source_sha256(source), {0: ids[::-1]})
    reopened = Presentation(output)
    assert [shape.shape_id for shape in reopened.slides[0].shapes] == ids[::-1]
    assert {shape.text for shape in reopened.slides[0].shapes} == {"Title", "Subtitle"}

    slide._element.append(OxmlElement("p:timing"))
    presentation.save(source)
    with pytest.raises(UnsupportedReadingOrder, match="animation"):
        validate_order(source, source_sha256(source), {0: ids[::-1]})


def test_media_shape_and_tampered_saved_package_are_refused(tmp_path):
    source, output, tampered = (
        tmp_path / "input.pptx",
        tmp_path / "output.pptx",
        tmp_path / "tampered.pptx",
    )
    ids = _deck(source)
    target = [ids[1], ids[0], ids[2]]
    repair_saved_order(source, output, source_sha256(source), {0: target})
    with ZipFile(output) as original, ZipFile(tampered, "w") as altered:
        for info in original.infolist():
            data = original.read(info.filename)
            if info.filename == "ppt/notesSlides/notesSlide1.xml":
                data = b"changed notes"
            altered.writestr(info, data)
    with pytest.raises(UnsupportedReadingOrder, match="Unrelated package part changed"):
        verify_saved_order(source, tampered, source_sha256(source), {0: target})

    presentation = Presentation(source)
    picture = presentation.slides[0].shapes[2]
    picture._element.append(OxmlElement("a:videoFile"))
    presentation.save(source)
    with pytest.raises(UnsupportedReadingOrder, match="media"):
        validate_order(source, source_sha256(source), {0: target})
