"""Recovered scans replace stale failure presentation only with durable success."""

from datetime import datetime, timedelta, timezone
import hashlib
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from src.db.models import (
    RemediationOutcome,
    ReviewAuditLog,
    Scan,
    ScanResult,
    ScanStatus,
    ScanType,
)
from src.education.remediation.output_claim import DescriptorBoundOutputClaim
from src.education.remediation.base import FixedIssue
from src.jobs import remediation_job
from src.jobs.remediation_subprocess import SubprocessRemediationResult
from src.services.remediation_artifact_service import ArtifactPublicationResult
from src.services.scan_fix_service import review_digest_for


@pytest.fixture
def context(tmp_path, monkeypatch):
    engine = create_engine("sqlite://")
    for model in (Scan, ScanResult, ReviewAuditLog):
        model.__table__.create(engine)
    source = tmp_path / "source.docx"
    document = Document()
    document.add_paragraph("Synthetic source")
    document.save(source)
    with Session(engine) as db:
        scan = Scan(
            id="scan-1",
            department_id="department-1",
            scan_type=ScanType.WORD,
            file_name=source.name,
            storage_path=str(source),
            status=ScanStatus.FAILED,
            remediation_outcome="manual_required",
            completed_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
            progress=42,
            progress_message="Remediation needs manual review",
            error_message="Previous attempt failed",
        )
        db.add(scan)
        db.add(ScanResult(scan_id=scan.id, compliance_score=80, issues=[]))
        db.commit()
        artifacts = MagicMock()
        artifacts.root = tmp_path / "managed"
        artifacts.claim_and_publish_stream.return_value = ArtifactPublicationResult(
            artifact=SimpleNamespace(
                id="artifact-1",
                mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                size_bytes=100,
                sha256="b" * 64,
                expires_at=datetime.now(timezone.utc) + timedelta(days=1),
                review_status="pending",
                provider_result={},
            ),
            artifact_id="artifact-1",
            publication_token="publication-1",
        )
        monkeypatch.setattr(
            remediation_job.RemediationArtifactService,
            "from_settings",
            lambda: artifacts,
        )
        monkeypatch.setattr(remediation_job, "persist_scan_fixes", MagicMock())
        notification = AsyncMock()
        monkeypatch.setattr(
            remediation_job, "_send_remediation_notification", notification
        )
        data = {
            "job_id": "job-1",
            "scan_id": scan.id,
            "department_id": scan.department_id,
            "provider": "local",
            "options": {"use_ai": False},
        }
        yield db, scan, data, artifacts, notification
    engine.dispose()


def _state(scan):
    return {
        key: getattr(scan, key)
        for key in (
            "status",
            "remediation_outcome",
            "completed_at",
            "progress",
            "progress_message",
            "error_message",
        )
    }


def _child(tmp_path, monkeypatch, db, scan, *, verified=True):
    result = db.query(ScanResult).filter(ScanResult.scan_id == scan.id).one()
    result.issues = [
        {
            "id": "issue-1",
            "category": "heading",
            "severity": "medium",
            "description": "Language is missing",
        }
    ]
    db.commit()
    candidate = tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("Synthetic verified output")
    document.save(candidate)
    claim = DescriptorBoundOutputClaim.from_path(
        candidate,
        display_path=str(candidate),
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    execution = SubprocessRemediationResult(
        values={
            "success": True,
            "verification_passed": verified,
            "total_issues": 1,
            "fixed_count": 1,
            "manual_count": 0,
            "failed_count": 0,
            "skipped_count": 0,
            "fixed_issues": [
                FixedIssue(
                    issue_id="issue-1",
                    category="heading",
                    severity="medium",
                    description="Language is missing",
                    fixed_content="en",
                    fix_method="rule",
                    needs_review=True,
                ).model_dump(mode="json")
            ],
            "manual_issues": [],
            "failed_issues": [],
        },
        output_claim=claim,
    )
    monkeypatch.setattr(
        remediation_job, "run_remediation_subprocess", AsyncMock(return_value=execution)
    )
    return claim


def _assert_complete(scan, outcome):
    assert scan.status == ScanStatus.COMPLETED
    assert scan.remediation_outcome == outcome
    assert scan.progress == 100
    assert scan.progress_message == "Remediation complete"
    assert scan.error_message is None


@pytest.mark.asyncio
async def test_zero_change_success_with_source_finding_stays_no_op(
    context, monkeypatch
):
    db, scan, data, artifacts, _ = context
    row = db.query(ScanResult).filter(ScanResult.scan_id == scan.id).one()
    row.issues = [{"id": "unfixed", "category": "heading"}]
    db.commit()
    child = SubprocessRemediationResult(
        values={
            "success": True,
            "verification_passed": True,
            "total_issues": 1,
            "fixed_count": 0,
            "manual_count": 0,
            "failed_count": 0,
            "skipped_count": 1,
            "fixed_issues": [],
            "manual_issues": [],
            "failed_issues": [],
        }
    )
    monkeypatch.setattr(
        remediation_job, "run_remediation_subprocess", AsyncMock(return_value=child)
    )

    response = await remediation_job.process_remediation_job(
        data, db, defer_final_commit=True, assert_owned=AsyncMock()
    )

    assert response["success"] is True
    assert scan.remediation_outcome == RemediationOutcome.NO_OP.value
    artifacts.claim_and_publish_stream.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize("binding", ["exact", "drift", "missing"])
async def test_normal_document_membership_uses_only_scanned_source_bytes(
    context, tmp_path, monkeypatch, binding
):
    db, scan, data, artifacts, _ = context
    source = tmp_path / "source.docx"
    scanned_bytes = source.read_bytes()
    scan.file_hash = (
        hashlib.sha256(scanned_bytes).hexdigest() if binding != "missing" else None
    )
    db.commit()
    if binding == "drift":
        source.write_bytes(scanned_bytes + b"changed since scan")
    claim = _child(tmp_path, monkeypatch, db, scan)
    fix = SimpleNamespace(occurrence_key="c" * 64)
    fix.review_digest = review_digest_for(fix)
    remediation_job.persist_scan_fixes.return_value = [fix]

    result = await remediation_job.process_remediation_job(
        data, db, defer_final_commit=True, assert_owned=AsyncMock()
    )

    assert result["success"] is True
    provider_result = (
        artifacts.claim_and_publish_stream.return_value.artifact.provider_result
    )
    membership = provider_result.get("reviewed_output_membership")
    if binding == "exact":
        assert membership["source_sha256"] == scan.file_hash
        assert membership["output_sha256"] == "b" * 64
        assert len(membership["fixes"]) == 1
    else:
        assert membership is None
    child_kwargs = remediation_job.run_remediation_subprocess.await_args.kwargs
    if binding == "missing":
        assert child_kwargs["source_path"] == str(source)
        assert "source_stream" not in child_kwargs
    else:
        assert child_kwargs["source_path"] == ""
        assert child_kwargs["source_stream"].getvalue() == source.read_bytes()
    assert claim.closed


@pytest.mark.asyncio
@pytest.mark.parametrize("has_issues", [False, True])
async def test_presentation_cleanup_is_staged_with_completion_and_rollback_restores_failure(
    context, tmp_path, monkeypatch, has_issues
):
    db, scan, data, artifacts, notification = context
    claim = _child(tmp_path, monkeypatch, db, scan) if has_issues else None
    before = _state(scan)
    result = await remediation_job.process_remediation_job(
        data, db, defer_final_commit=True, assert_owned=AsyncMock()
    )
    assert result["success"] is True
    _assert_complete(scan, "completed" if has_issues else "no_op")
    assert db.in_transaction()
    notification.assert_not_awaited()
    db.rollback()
    assert _state(scan) == before
    if has_issues:
        artifacts.claim_and_publish_stream.assert_called_once()
        assert claim.closed
    else:
        artifacts.claim_and_publish_stream.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize("has_issues", [False, True])
async def test_success_commit_retains_clean_presentation(
    context, tmp_path, monkeypatch, has_issues
):
    db, scan, data, _, _ = context
    if has_issues:
        _child(tmp_path, monkeypatch, db, scan)
    result = await remediation_job.process_remediation_job(
        data, db, defer_final_commit=True, assert_owned=AsyncMock()
    )
    assert result["success"] is True
    db.commit()
    db.expire_all()
    _assert_complete(scan, "completed" if has_issues else "no_op")


@pytest.mark.asyncio
async def test_failed_verification_keeps_old_scan_presentation(
    context, tmp_path, monkeypatch
):
    db, scan, data, artifacts, _ = context
    _child(tmp_path, monkeypatch, db, scan, verified=False)
    before = _state(scan)
    result = await remediation_job.process_remediation_job(
        data, db, defer_final_commit=True, assert_owned=AsyncMock()
    )
    assert result["success"] is False
    assert _state(scan) == before
    artifacts.claim_and_publish_stream.assert_not_called()


@pytest.mark.asyncio
async def test_post_publication_failure_restores_all_presentation_fields(
    context, tmp_path, monkeypatch
):
    db, scan, data, artifacts, _ = context
    _child(tmp_path, monkeypatch, db, scan)
    before = _state(scan)
    monkeypatch.setattr(
        db, "add", MagicMock(side_effect=RuntimeError("audit write failed"))
    )
    with pytest.raises(remediation_job.RetryableRemediationJobError):
        await remediation_job.process_remediation_job(
            data, db, defer_final_commit=True, assert_owned=AsyncMock()
        )
    assert _state(scan) == before
    artifacts.abort_staging.assert_called_once_with(
        db, artifact_id="artifact-1", publication_token="publication-1"
    )


@pytest.mark.asyncio
async def test_noop_commit_failure_restores_all_presentation_fields(
    context, monkeypatch
):
    db, scan, data, artifacts, _ = context
    before = _state(scan)
    monkeypatch.setattr(
        db, "commit", MagicMock(side_effect=RuntimeError("commit failed"))
    )
    result = await remediation_job.process_remediation_job(data, db)
    assert result["success"] is False
    assert _state(scan) == before
    artifacts.claim_and_publish_stream.assert_not_called()
