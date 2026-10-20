"""Image descriptions must never be written without the image's pixels.

When vision analysis of an image is unavailable or unsuccessful, the PDF, DOCX
and PPTX remediators must leave the finding open for human review. Describing
an image from the text around it produces content that was never checked
against the image, even when a checker would accept it.
"""

import io
import zipfile
from pathlib import Path

import pikepdf
import pytest
from docx import Document
from PIL import Image
from pptx import Presentation

from src.education.docx_processor import DocxProcessor
from src.education.pdf_processor import PDFProcessor
from src.education.pptx_processor import PowerPointProcessor
from src.education.remediation.base import IssueCategory, RemediationConfig
from src.education.remediation.docx_remediator import DocxRemediator
from src.education.remediation.pdf_remediator import PdfRemediator
from src.education.remediation.pptx_remediator import PptxRemediator
from tests.test_pdf_existing_figure_alt import _fixture as _make_tagged_pdf

pytestmark = pytest.mark.unit

INVENTED = "A plausible description"


class _TextOnlyClient:
    """Purpose-bound client with no vision method but a working text method."""

    def __init__(self):
        self.text_calls = 0
        self.vision_calls = 0

    def generate_text_sync(self, *args, **kwargs):
        self.text_calls += 1
        return {"success": True, "content": INVENTED}


class _VisionFails(_TextOnlyClient):
    def analyze_image_sync(self, *args, **kwargs):
        self.vision_calls += 1
        return {"success": False, "content": "", "error": "provider unavailable"}


class _VisionRaises(_TextOnlyClient):
    def analyze_image_sync(self, *args, **kwargs):
        self.vision_calls += 1
        raise RuntimeError("provider unavailable")


def _client(mode):
    return {"fails": _VisionFails, "raises": _VisionRaises, "absent": _TextOnlyClient}[
        mode
    ]()


def _png() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (24, 24), "green").save(buffer, format="PNG")
    return buffer.getvalue()


def _config(tmp_path):
    return RemediationConfig(
        use_ai=True,
        allow_legacy_nested_ai=False,
        create_backup=False,
        output_directory=str(tmp_path / "output"),
    )


def _alt_issues(issues):
    return [
        issue
        for issue in issues
        if issue.get("issue_type") == "missing_alt_text"
        or (issue.get("category") == "alt_text" and "missing" in str(issue))
    ]


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("A lecture on cell biology and the mitochondrion.")
    document.add_picture(io.BytesIO(_png()))
    document.save(path)


def _make_pptx(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Cell biology"
    picture = slide.shapes.add_picture(io.BytesIO(_png()), 0, 0)
    # python-pptx fills descr with the file name; leave the picture undescribed.
    picture._element.nvPicPr.cNvPr.attrib.pop("descr", None)
    presentation.save(path)


def _office_text(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        return "".join(
            archive.read(name).decode("utf-8", "replace")
            for name in archive.namelist()
            if name.endswith(".xml")
        )


def _pdf_contains(path: str, needle: str) -> bool:
    with pikepdf.open(path) as pdf:
        for obj in pdf.objects:
            if isinstance(obj, pikepdf.Dictionary) and "/Alt" in obj:
                if needle in str(obj["/Alt"]):
                    return True
    return needle.encode("utf-16-be") in Path(path).read_bytes() or (
        needle.encode() in Path(path).read_bytes()
    )


def _assert_failed_closed(result, client, alt_text_in_output):
    alt_fixed = [
        fix for fix in result.fixed_issues if fix.category == IssueCategory.ALT_TEXT
    ]
    alt_manual = [
        issue
        for issue in result.manual_issues
        if issue.category == IssueCategory.ALT_TEXT
    ]
    assert alt_fixed == [], f"alt text reported fixed without pixels: {alt_fixed}"
    assert alt_manual, "image issue was not left open for human review"
    assert not alt_text_in_output, "invented description was written to the document"
    assert client.text_calls == 0, "text-only generation was attempted for an image"


@pytest.mark.parametrize("vision", ["fails", "raises", "absent"])
def test_docx_never_describes_an_image_without_its_pixels(tmp_path, vision):
    source = tmp_path / "source.docx"
    _make_docx(source)
    issues = _alt_issues(DocxProcessor().process_docx(str(source)).issues)
    assert issues, "scanner reported no missing alt text for the fixture"

    client = _client(vision)
    result = DocxRemediator(
        str(source), issues, _config(tmp_path), alt_text_client=client
    ).remediate()

    written = bool(result.output_file) and INVENTED in _office_text(result.output_file)
    _assert_failed_closed(result, client, written)


@pytest.mark.parametrize("vision", ["fails", "raises", "absent"])
def test_pptx_never_describes_an_image_without_its_pixels(tmp_path, vision):
    source = tmp_path / "source.pptx"
    _make_pptx(source)
    issues = _alt_issues(PowerPointProcessor().process_pptx(str(source)).issues)
    assert issues, "scanner reported no missing alt text for the fixture"

    client = _client(vision)
    result = PptxRemediator(
        str(source), issues, _config(tmp_path), alt_text_client=client
    ).remediate()

    written = bool(result.output_file) and INVENTED in _office_text(result.output_file)
    _assert_failed_closed(result, client, written)


@pytest.mark.parametrize("vision", ["fails", "raises", "absent"])
def test_pdf_never_describes_an_image_without_its_pixels(tmp_path, vision):
    source = tmp_path / "source.pdf"
    _make_tagged_pdf(source)
    issues = _alt_issues(PDFProcessor().process_pdf(str(source)).issues)
    assert issues, "scanner reported no missing alt text for the fixture"

    client = _client(vision)
    result = PdfRemediator(
        str(source), issues, _config(tmp_path), alt_text_client=client
    ).remediate()
    try:
        written = bool(result.output_file) and _pdf_contains(
            result.output_file, INVENTED
        )
        _assert_failed_closed(result, client, written)
    finally:
        result.close_output_claim()
