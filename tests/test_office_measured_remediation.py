"""Saved Office documents must determine scores and verified fix counts."""

import hashlib
import pytest
from io import BytesIO
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from pptx import Presentation
from PIL import Image

from src.education.docx_processor import DocxProcessor
from src.education.pptx_processor import PowerPointProcessor
from src.education.xlsx_processor import XlsxProcessor
from src.education.remediation.base import RemediationConfig
from src.education.remediation.docx_remediator import DocxRemediator
from src.education.remediation.pptx_remediator import PptxRemediator
from src.education.remediation.xlsx_remediator import XlsxRemediator
from src.education.compliance_scoring import calculate_compliance_score


def office_fixture(tmp_path, kind):
    path = tmp_path / f"source.{kind}"
    if kind == "docx":
        document = Document()
        document.add_heading("Course overview", level=2)
        document.add_paragraph("Tiny body text").runs[0].font.size = Pt(8)
        document.save(path)
        return path, DocxProcessor().process_docx, DocxRemediator
    if kind == "pptx":
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[0])
        presentation.save(path)
        return path, PowerPointProcessor().process_pptx, PptxRemediator
    workbook = Workbook()
    workbook.active.append(["Course", "Credits"])
    workbook.active.append(["History", 3])
    workbook.save(path)
    return path, XlsxProcessor().process_xlsx, XlsxRemediator


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_scores_equal_independent_source_and_saved_scans(tmp_path, kind):
    source, scan, remediator_type = office_fixture(tmp_path, kind)
    original_bytes = source.read_bytes()
    before = scan(str(source))
    assert before.issues, "Fixture must exercise actual scanner findings"
    result = remediator_type(
        str(source), before.issues, RemediationConfig(create_backup=False)
    ).remediate()
    assert result.success, result.error_message
    after = scan(result.output_file)
    assert result.original_compliance_score == before.compliance_score
    assert result.remediated_compliance_score == after.compliance_score
    assert result.improvement == after.compliance_score - before.compliance_score
    assert result.score_provenance == "scanner_rescan"
    assert (
        result.score_measurement["source_sha256"]
        == hashlib.sha256(source.read_bytes()).hexdigest()
    )
    from pathlib import Path

    assert (
        result.score_measurement["output_sha256"]
        == hashlib.sha256(Path(result.output_file).read_bytes()).hexdigest()
    )
    assert result.score_verification_reason is None
    assert source.read_bytes() == original_bytes
    assert result.verification_result.issues_after == len(after.issues)
    assert result.fixed_count == len(result.fixed_issues)
    assert all(fix.verification_passed for fix in result.fixed_issues)


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_successful_write_that_does_not_fix_finding_is_manual(
    tmp_path, monkeypatch, kind
):
    source, scan, remediator_type = office_fixture(tmp_path, kind)
    before = scan(str(source))
    remediator = remediator_type(
        str(source), before.issues, RemediationConfig(create_backup=False)
    )
    monkeypatch.setattr(remediator, "apply_fix", lambda *args: True)
    monkeypatch.setattr(remediator, "can_auto_fix", lambda *args: True)
    monkeypatch.setattr(
        remediator, "_generate_fix", lambda *args: "An attempted change"
    )
    result = remediator.remediate()
    assert result.success, result.error_message
    assert result.fixed_count == 0
    assert not result.fixed_issues
    assert result.manual_count > 0
    assert not result.verification_passed
    assert result.remediated_compliance_score == before.compliance_score


def test_real_docx_heading_fix_at_paragraph_zero_is_verified(tmp_path):
    source, scan, remediator_type = office_fixture(tmp_path, "docx")
    before = scan(str(source))
    heading = next(issue for issue in before.issues if issue["category"] == "heading")
    result = remediator_type(
        str(source), [heading], RemediationConfig(create_backup=False)
    ).remediate()
    assert result.success, result.error_message
    assert Document(result.output_file).paragraphs[0].style.name == "Heading 1"
    assert result.fixed_count == 1
    assert result.fixed_issues[0].verification_passed
    assert (
        result.remediated_compliance_score == scan(result.output_file).compliance_score
    )


def test_disabled_verification_never_reports_estimated_success(tmp_path):
    source, scan, remediator_type = office_fixture(tmp_path, "docx")
    result = remediator_type(
        str(source),
        scan(str(source)).issues,
        RemediationConfig(create_backup=False, verify_fixes=False),
    ).remediate()
    assert result.success
    assert result.remediated_compliance_score is None
    assert result.improvement is None
    assert result.fixed_count == 0
    assert result.manual_count > 0
    assert not result.verification_passed


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
def test_scanner_score_uses_reported_severities_and_critical_cap(tmp_path, kind):
    source, scan, _ = office_fixture(tmp_path, kind)
    image = BytesIO()
    Image.new("RGB", (20, 20), color="red").save(image, format="PNG")
    image.seek(0)
    if kind == "docx":
        document = Document(source)
        document.add_picture(image)
        document.save(source)
    elif kind == "pptx":
        presentation = Presentation(source)
        presentation.slides[0].shapes.add_picture(image, 0, 0)
        presentation.save(source)
    else:
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as SheetImage

        workbook = load_workbook(source)
        workbook.active.add_image(SheetImage(image), "D1")
        workbook.save(source)
    result = scan(str(source))
    assert any(issue["severity"] == "critical" for issue in result.issues)
    assert result.compliance_score <= 49
    if kind == "docx":
        count = sum(
            getattr(result, key)
            for key in (
                "total_paragraphs",
                "total_images",
                "total_tables",
                "total_lists",
                "total_links",
            )
        )
    elif kind == "pptx":
        count = result.total_shapes + result.total_images
    else:
        count = (
            result.total_rows
            + result.total_charts
            + result.total_images
            + result.total_sheets
        )
    assert (
        result.compliance_score
        == calculate_compliance_score(result.issues, count).score
    )


def test_output_scanner_failure_is_unavailable_and_manual(tmp_path, monkeypatch):
    from src.education.remediation import office_verification

    source, scan, remediator_type = office_fixture(tmp_path, "docx")
    before = scan(str(source))
    real_scan = office_verification.scan_office

    def fail_on_output(document_type, path):
        if path != str(source):
            raise RuntimeError("Scanner unavailable")
        return real_scan(document_type, path)

    monkeypatch.setattr(office_verification, "scan_office", fail_on_output)
    result = remediator_type(
        str(source), before.issues, RemediationConfig(create_backup=False)
    ).remediate()
    assert result.success
    assert result.original_compliance_score == before.compliance_score
    assert result.remediated_compliance_score is None
    assert result.score_provenance is None
    assert result.fixed_count == 0
    assert not result.verification_passed
    assert result.manual_count > 0


def test_semantic_findings_cannot_be_verified_by_presence(tmp_path):
    source, scan, remediator_type = office_fixture(tmp_path, "docx")
    issues = scan(str(source)).issues
    issues.append(
        {
            "id": "semantic_alt",
            "category": "alt_text",
            "severity": "high",
            "description": "Existing alternative text is inaccurate",
            "location": "Image 1",
            "alt_text_validated": True,
        }
    )
    result = remediator_type(
        str(source), issues, RemediationConfig(create_backup=False, use_ai=False)
    ).remediate()
    assert result.remediated_compliance_score is None
    assert result.score_provenance is None
    assert result.fixed_count == 0
    assert not result.verification_passed


def test_saved_output_regression_is_measured_and_flagged(tmp_path, monkeypatch):
    source, scan, remediator_type = office_fixture(tmp_path, "docx")
    before = scan(str(source))
    heading = next(issue for issue in before.issues if issue["category"] == "heading")
    remediator = remediator_type(
        str(source), [heading], RemediationConfig(create_backup=False)
    )
    save = remediator._save_document

    def save_with_regression(document):
        image = BytesIO()
        Image.new("RGB", (20, 20), color="red").save(image, format="PNG")
        image.seek(0)
        document.add_picture(image)
        return save(document)

    monkeypatch.setattr(remediator, "_save_document", save_with_regression)
    result = remediator.remediate()
    assert (
        result.remediated_compliance_score == scan(result.output_file).compliance_score
    )
    assert result.improvement < 0
    assert result.fixed_count == 1
    assert result.verification_result.regressions
    assert not result.verification_passed


@pytest.mark.parametrize("kind", ["docx", "pptx", "xlsx"])
@pytest.mark.parametrize("alias", ["direct", "symlink", "hardlink"])
def test_original_cannot_be_used_as_output(tmp_path, kind, alias):
    source, scan, remediator_type = office_fixture(tmp_path, kind)
    original = source.read_bytes()
    output = source
    if alias != "direct":
        output = tmp_path / f"alias.{kind}"
        if alias == "symlink":
            output.symlink_to(source)
        else:
            output.hardlink_to(source)
    result = remediator_type(
        str(source),
        scan(str(source)).issues,
        RemediationConfig(create_backup=False, output_filename=output.name),
    ).remediate()
    assert not result.success
    assert "must not overwrite" in result.error_message
    assert source.read_bytes() == original


def test_pptx_animation_and_embedded_media_findings_affect_score(tmp_path, monkeypatch):
    from src.education.pptx_processor import AnimationIssue, EmbeddedMediaIssue

    source, _, _ = office_fixture(tmp_path, "pptx")
    processor = PowerPointProcessor()
    monkeypatch.setattr(
        processor,
        "_analyze_animations",
        lambda *args: [
            AnimationIssue(
                slide_number=1,
                animation_index=0,
                animation_type="emphasis",
                element_name="Flash",
                duration_ms=100,
                issues=["rapid_flash"],
                suggested_fix="Remove rapid flashing",
            )
        ],
    )
    monkeypatch.setattr(
        processor,
        "_check_embedded_media",
        lambda *args: [
            EmbeddedMediaIssue(
                slide_number=1,
                media_index=0,
                media_type="video",
                issue_type="missing_captions",
                recommendations=["Add captions"],
                suggested_fix="Add captions",
            )
        ],
    )
    result = processor.process_pptx(str(source))
    assert any(issue["id"].startswith("animation_") for issue in result.issues)
    assert any(issue["id"].startswith("embedded_media_") for issue in result.issues)
    assert result.summary["total_issues"] == len(result.issues)
    assert result.compliance_score <= 49
