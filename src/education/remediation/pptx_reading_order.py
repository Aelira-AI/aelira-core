"""Conservative, source-bound repair of a saved PPTX shape-tree order.

PresentationML uses the lexical order of top-level ``p:spTree`` objects for
both navigation and visual stacking.  Only an explicit complete permutation
of stable shape IDs is accepted, and inverted overlapping pairs are refused.
The ZIP writer copies every other package member without changing its data.
"""

from __future__ import annotations

import hashlib
import os
import shutil
import tempfile
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
SHAPES = {f"{{{P}}}sp", f"{{{P}}}pic"}
MAX_SHAPES = 128
MAX_TARGET_SLIDES = 64
UNSAFE_DESCENDANTS = {
    f"{{{A}}}effectLst",
    f"{{{A}}}effectDag",
    f"{{{A}}}scene3d",
    f"{{{A}}}sp3d",
    f"{{{A}}}videoFile",
    f"{{{A}}}audioFile",
}


class UnsupportedReadingOrder(ValueError):
    """The requested ordering cannot be proved safe for this presentation."""


def source_sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _shape_ids(tree):
    children = list(tree)
    if len(children) < 2 or [node.tag for node in children[:2]] != [
        f"{{{P}}}nvGrpSpPr",
        f"{{{P}}}grpSpPr",
    ]:
        raise UnsupportedReadingOrder("Unexpected slide shape-tree structure")
    if any(node.tag not in SHAPES for node in children[2:]):
        raise UnsupportedReadingOrder("Grouped or unsupported slide object")
    if not 1 <= len(children[2:]) <= MAX_SHAPES:
        raise UnsupportedReadingOrder("Slide shape count outside bounded repair")
    ids = []
    for node in children[2:]:
        props = node.find(f"./{{{P}}}nvSpPr/{{{P}}}cNvPr")
        if props is None:
            props = node.find(f"./{{{P}}}nvPicPr/{{{P}}}cNvPr")
        if props is None or not props.get("id", "").isdigit():
            raise UnsupportedReadingOrder("Shape has no stable numeric ID")
        ids.append(int(props.get("id")))
        if props.find(f"./{{{A}}}extLst") is not None or any(
            etree.QName(desc).localname.lower() == "decorative" for desc in node.iter()
        ):
            raise UnsupportedReadingOrder("Decorative or extended shape reading state")
        if any(desc.tag in UNSAFE_DESCENDANTS for desc in node.iter()):
            raise UnsupportedReadingOrder(
                "Shape uses unsupported visual or media effects"
            )
        for transform in node.iter(f"{{{A}}}xfrm"):
            if transform.get("rot") not in (None, "0") or any(
                transform.get(flag) in ("1", "true") for flag in ("flipH", "flipV")
            ):
                raise UnsupportedReadingOrder("Rotated or flipped shape geometry")
    if len(set(ids)) != len(ids):
        raise UnsupportedReadingOrder("Duplicate shape IDs")
    return ids


def _tree(root):
    tree = root.find(f"./{{{P}}}cSld/{{{P}}}spTree")
    if tree is None or root.find(f"./{{{P}}}timing") is not None:
        raise UnsupportedReadingOrder(
            "Slide has no supported shape tree or has animation"
        )
    return tree


def _overlap(first, second, slide_width, slide_height):
    for shape in (first, second):
        if any(
            not isinstance(value, int) or isinstance(value, bool)
            for value in (shape.left, shape.top, shape.width, shape.height)
        ):
            raise UnsupportedReadingOrder("Unknown shape geometry")
        if shape.width <= 0 or shape.height <= 0:
            raise UnsupportedReadingOrder("Unknown shape bounds")
        if (
            shape.left < 0
            or shape.top < 0
            or shape.left + shape.width > slide_width
            or shape.top + shape.height > slide_height
        ):
            raise UnsupportedReadingOrder("Shape extends outside slide bounds")
    return (
        first.left <= second.left + second.width
        and second.left <= first.left + first.width
        and first.top <= second.top + second.height
        and second.top <= first.top + first.height
    )


def _slide_info(presentation, slide_index, requested):
    if type(slide_index) is not int or not 0 <= slide_index < len(presentation.slides):
        raise UnsupportedReadingOrder("Invalid slide index")
    slide = presentation.slides[slide_index]
    part_name = str(slide.part.partname).lstrip("/")
    if (
        not isinstance(requested, list)
        or not requested
        or any(type(item) is not int or item <= 0 for item in requested)
    ):
        raise UnsupportedReadingOrder("Expected a complete numeric shape-ID order")
    root = etree.fromstring(slide.part.blob)
    current = _shape_ids(_tree(root))
    if (
        len(requested) != len(current)
        or set(requested) != set(current)
        or len(set(requested)) != len(requested)
    ):
        raise UnsupportedReadingOrder(
            "Target does not cover each source shape ID exactly once"
        )
    by_id = {shape.shape_id: shape for shape in slide.shapes}
    if set(by_id) != set(current):
        raise UnsupportedReadingOrder("Shape identity could not be resolved")
    new_rank = {value: index for index, value in enumerate(requested)}
    for first_index, first_id in enumerate(current):
        for second_id in current[first_index + 1 :]:
            if new_rank[first_id] > new_rank[second_id] and _overlap(
                by_id[first_id],
                by_id[second_id],
                presentation.slide_width,
                presentation.slide_height,
            ):
                raise UnsupportedReadingOrder(
                    "Order would change overlapping visual stacking"
                )
    return part_name, current


def validate_order(
    source: str | Path, expected_sha256: str, targets: dict[int, list[int]]
):
    """Validate reviewer order and return slide part names with original IDs."""
    if not isinstance(expected_sha256, str) or source_sha256(source) != expected_sha256:
        raise UnsupportedReadingOrder("Source hash does not match accepted order")
    if not isinstance(targets, dict) or not 1 <= len(targets) <= MAX_TARGET_SLIDES:
        raise UnsupportedReadingOrder("No accepted slide order")
    presentation = Presentation(str(source))
    specs = {}
    for index, requested in targets.items():
        part_name, current = _slide_info(presentation, index, requested)
        specs[part_name] = (index, current, requested)
    if len(specs) != len(targets):
        raise UnsupportedReadingOrder("Slide part identity is ambiguous")
    return specs


def _reordered_xml(source_xml: bytes, target: list[int]) -> bytes:
    root = etree.fromstring(source_xml)
    tree = _tree(root)
    current = _shape_ids(tree)
    by_id = dict(zip(current, list(tree)[2:]))
    for node in list(tree)[2:]:
        tree.remove(node)
    for shape_id in target:
        tree.append(by_id[shape_id])
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _read_parts(path):
    with ZipFile(path) as archive:
        names = archive.namelist()
        if len(set(names)) != len(names):
            raise UnsupportedReadingOrder("Duplicate package member names")
        if any(name.startswith("_xmlsignatures/") for name in names):
            raise UnsupportedReadingOrder("Digitally signed presentation")
        return {name: archive.read(name) for name in names}


def _verify_saved(source, candidate, specs):
    original = _read_parts(source)
    saved = _read_parts(candidate)
    if set(original) != set(saved):
        raise UnsupportedReadingOrder("Saved package member set changed")
    for name, data in original.items():
        if name not in specs:
            if saved[name] != data:
                raise UnsupportedReadingOrder(f"Unrelated package part changed: {name}")
            continue
        _, old_order, target = specs[name]
        result = etree.fromstring(saved[name])
        result_tree = _tree(result)
        if _shape_ids(result_tree) != target:
            raise UnsupportedReadingOrder(
                "Saved shape order differs from accepted order"
            )
        saved_shapes = dict(zip(target, list(result_tree)[2:]))
        for node in list(result_tree)[2:]:
            result_tree.remove(node)
        for shape_id in old_order:
            result_tree.append(saved_shapes[shape_id])
        if etree.tostring(result, method="c14n") != etree.tostring(
            etree.fromstring(data), method="c14n"
        ):
            raise UnsupportedReadingOrder(
                "Slide content changed outside the shape order"
            )
    reopened = Presentation(str(candidate))
    for _, (index, _, target) in specs.items():
        if [shape.shape_id for shape in reopened.slides[index].shapes] != target:
            raise UnsupportedReadingOrder(
                "Reopened presentation order differs from target"
            )


def verify_saved_order(
    source: str | Path,
    output: str | Path,
    expected_sha256: str,
    targets: dict[int, list[int]],
) -> None:
    """Recheck the exact delivered file against its source and accepted order."""
    specs = validate_order(source, expected_sha256, targets)
    _verify_saved(source, output, specs)


def repair_saved_order(
    source: str | Path,
    output: str | Path,
    expected_sha256: str,
    targets: dict[int, list[int]],
) -> None:
    """Atomically publish only after package and reopened-order verification."""
    source, output = Path(source), Path(output)
    if source.resolve() == output.resolve() or (
        output.exists() and source.samefile(output)
    ):
        raise UnsupportedReadingOrder("Output must differ from source")
    output.parent.mkdir(parents=True, exist_ok=True)
    snapshot_fd, snapshot = tempfile.mkstemp(
        prefix=".pptx-source-", suffix=".pptx", dir=output.parent
    )
    os.close(snapshot_fd)
    descriptor, candidate = tempfile.mkstemp(
        prefix=".pptx-order-", suffix=".pptx", dir=output.parent
    )
    os.close(descriptor)
    try:
        # Bind validation, mutation and preservation checks to the same bytes.
        # Reopening the caller's path at each stage permits a replaced source
        # to pass preservation checks against itself after hash validation.
        shutil.copyfile(source, snapshot)
        specs = validate_order(snapshot, expected_sha256, targets)
        with ZipFile(snapshot) as original, ZipFile(candidate, "w") as saved:
            for info in original.infolist():
                data = original.read(info.filename)
                if info.filename in specs:
                    data = _reordered_xml(data, specs[info.filename][2])
                saved.writestr(info, data)
        _verify_saved(snapshot, candidate, specs)
        if source_sha256(source) != expected_sha256:
            raise UnsupportedReadingOrder("Source changed during repair")
        os.replace(candidate, output)
    finally:
        if os.path.exists(candidate):
            os.unlink(candidate)
        os.unlink(snapshot)
