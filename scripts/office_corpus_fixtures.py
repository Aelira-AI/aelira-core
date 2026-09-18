"""Repository-authored, CC0 synthetic Office documents; no external assets."""

from datetime import datetime
from io import BytesIO
from pathlib import Path
import wave
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as DOC_REL
from lxml import etree
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

FIXED_TIME = datetime(2000, 1, 1)


def image_bytes():
    stream = BytesIO()
    Image.new("RGB", (24, 24), "#276090").save(stream, format="PNG")
    stream.seek(0)
    return stream


def _metadata(properties):
    properties.created = FIXED_TIME
    properties.modified = FIXED_TIME
    if hasattr(properties, "author"):
        properties.author = "Synthetic corpus"
    if hasattr(properties, "creator"):
        properties.creator = "Synthetic corpus"


def _normalize_zip(path):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    # openpyxl replaces modified during save, after our metadata assignment.
    # Freeze serialized timestamps, not just the in-memory properties.
    core = etree.fromstring(parts["docProps/core.xml"])
    for element in core:
        if element.tag in {
            "{http://purl.org/dc/terms/}created",
            "{http://purl.org/dc/terms/}modified",
        }:
            element.text = "2000-01-01T00:00:00Z"
    parts["docProps/core.xml"] = etree.tostring(core)
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, payload in sorted(parts.items()):
            info = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, payload)


def _docx(path, corrected):
    document = Document()
    _metadata(document.core_properties)
    document.core_properties.title = "Course workbook" if corrected else ""
    document.add_heading("Course workbook", 1)
    document.add_paragraph("Preserve café, equations x + y = 3, and teaching notes.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, (("Topic", "Value"), ("History", "3"))):
        for cell, text in zip(row.cells, values):
            cell.text = text
    document.sections[0].header.paragraphs[0].text = "Course header"
    document.sections[0].footer.paragraphs[0].text = "Course footer"
    document.add_picture(image_bytes())
    paragraph = document.add_paragraph()
    link = OxmlElement("w:hyperlink")
    link.set(
        qn("r:id"),
        document.part.relate_to(
            "https://example.org/course", DOC_REL.HYPERLINK, is_external=True
        ),
    )
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Course reference"
    run.append(text)
    link.append(run)
    paragraph._p.append(link)
    document.save(path)


def _pptx(path, corrected):
    presentation = Presentation()
    _metadata(presentation.core_properties)
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Course overview"
    shape = slide.shapes.title
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
    run = shape.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(32)
    run.font.color.rgb = RGBColor(*((105, 105, 105) if corrected else (150, 150, 150)))
    run.hyperlink.address = "https://example.org/course"
    slide.shapes.add_picture(image_bytes(), Inches(1), Inches(2))
    slide.notes_slide.notes_text_frame.text = "Speaker notes and synthetic transcript."
    audio = BytesIO()
    with wave.open(audio, "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(2)
        handle.setframerate(8000)
        handle.writeframes(b"\x00\x00" * 80)
    audio.seek(0)
    # Real WAV bytes, plus a generated poster. PowerPoint playback is not claimed.
    slide.shapes.add_movie(
        audio,
        Inches(3),
        Inches(2),
        Inches(1),
        Inches(1),
        poster_frame_image=image_bytes(),
        mime_type="audio/wav",
    )
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "Unchanged second slide"
    second.notes_slide.notes_text_frame.text = "Second-slide notes."
    # Preserve timing markup, without claiming to exercise animation playback.
    ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    timing = etree.SubElement(second._element, f"{{{ns}}}timing")
    sequence = etree.SubElement(timing, f"{{{ns}}}tnLst")
    parallel = etree.SubElement(sequence, f"{{{ns}}}par")
    etree.SubElement(
        parallel, f"{{{ns}}}cTn", id="1", dur="indefinite", nodeType="tmRoot"
    )
    presentation.save(path)


def _xlsx(path, corrected):
    workbook = Workbook()
    _metadata(workbook.properties)
    sheet = workbook.active
    sheet.title = "Course results"
    sheet.append(["Student", "Marks", "Adjusted", "Notes"])
    for row in range(2, 32):
        sheet.append([f"Synthetic {row}", row, f"=B{row}*2", "Teaching"])
    sheet["A2"].hyperlink = "https://example.org/course"
    sheet["A2"].comment = Comment("Teaching note", "Synthetic corpus")
    sheet.auto_filter.ref = "A1:C31"
    sheet.freeze_panes = "A2" if corrected else None
    sheet.print_area = "A1:C31"
    workbook.defined_names.add(
        DefinedName("Marks", attr_text="'Course results'!$B$2:$B$31")
    )
    chart = BarChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=5), titles_from_data=True
    )
    sheet.add_chart(chart, "E2")
    workbook.create_sheet("Unchanged")["A1"] = "Preserve this worksheet"
    workbook.save(path)


def generate_sources(root: Path, *, corrected=False):
    """Never overwrite an existing corpus directory."""
    root.mkdir(parents=True, exist_ok=False)
    paths = {}
    for kind, generate in (("docx", _docx), ("pptx", _pptx), ("xlsx", _xlsx)):
        path = root / f"course.{kind}"
        generate(path, corrected)
        _normalize_zip(path)
        paths[kind] = path
    return paths
