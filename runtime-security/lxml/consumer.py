"""Synthetic local XML and real pinned Office-library controls as a nonroot user."""

from __future__ import annotations

import io
import json
import runpy
import tempfile
from pathlib import Path

from lxml import etree
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook, load_workbook


def xml_controls() -> None:
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        (root / "local.dtd").write_text(
            "<!ELEMENT campus (name)><!ELEMENT name (#PCDATA)>"
        )
        (root / "catalog.xml").write_text(
            '<?xml version="1.0"?><catalog xmlns="urn:oasis:names:tc:entity:xmlns:xml:catalog"><public publicId="-//Aelira//Campus DTD//EN" uri="local.dtd"/></catalog>'
        )
        (root / "document.xml").write_text(
            '<!DOCTYPE campus PUBLIC "-//Aelira//Campus DTD//EN" "file:///absent.dtd"><campus><name>Accessible</name></campus>'
        )
        parser = etree.XMLParser(load_dtd=True, dtd_validation=True, no_network=True)
        import os

        old = os.environ.get("XML_CATALOG_FILES")
        os.environ["XML_CATALOG_FILES"] = str(root / "catalog.xml")
        try:
            tree = etree.parse(str(root / "document.xml"), parser)
        finally:
            if old is None:
                os.environ.pop("XML_CATALOG_FILES", None)
            else:
                os.environ["XML_CATALOG_FILES"] = old
        assert tree.xpath("string(/campus/name)") == "Accessible"
        assert etree.tostring(tree, method="c14n").endswith(b"</campus>")
        stylesheet = etree.XML(
            b'<xsl:stylesheet xmlns:xsl="http://www.w3.org/1999/XSL/Transform" version="1.0"><xsl:template match="/"><out><xsl:value-of select="/campus/name"/></out></xsl:template></xsl:stylesheet>'
        )
        assert b"Accessible" in bytes(etree.XSLT(stylesheet)(tree))
        rng = etree.RelaxNG(
            etree.XML(
                b'<element xmlns="http://relaxng.org/ns/structure/1.0" name="campus"><element name="name"><text/></element></element>'
            )
        )
        assert rng.validate(tree)
        (root / "part.xml").write_text("<name>Included</name>")
        part = root / "part.xml"
        included = etree.ElementTree(
            etree.XML(
                f'<campus xmlns:xi="http://www.w3.org/2001/XInclude"><xi:include href="{part}"/></campus>'.encode()
            )
        )
        included.xinclude()
        assert included.xpath("string(/campus/name)") == "Included"
        for payload in (b"<campus><name>", b"<campus>" + b"<" * 100):
            try:
                etree.fromstring(payload)
            except etree.XMLSyntaxError:
                pass
            else:
                raise AssertionError("malformed XML accepted")


def office_controls() -> None:
    doc = Document()
    doc.add_heading("Accessible campus", 1)
    doc.add_paragraph("Synthetic content")
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    assert Document(buffer).paragraphs[1].text == "Synthetic content"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = "Slide content"
    buffer = io.BytesIO()
    deck.save(buffer)
    buffer.seek(0)
    assert Presentation(buffer).slides[0].shapes[0].text == "Slide content"
    # Exercise the exact public application's reading-order validator source.
    source = Path("/inputs/source/src/education/remediation/pptx_reading_order.py")
    assert source.is_file()
    app = runpy.run_path(str(source))
    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "deck.pptx"
        path.write_bytes(buffer.getvalue())
        shape_id = Presentation(str(path)).slides[0].shapes[0].shape_id
        assert app["validate_order"](path, app["source_sha256"](path), {0: [shape_id]})
    book = Workbook()
    book.active["A1"] = "Spreadsheet content"
    buffer = io.BytesIO()
    book.save(buffer)
    buffer.seek(0)
    assert load_workbook(buffer).active["A1"].value == "Spreadsheet content"


if __name__ == "__main__":
    xml_controls()
    office_controls()
    print(
        json.dumps(
            {
                "nonroot": True,
                "xml": [
                    "catalog-dtd",
                    "xpath",
                    "xslt",
                    "relaxng",
                    "xinclude",
                    "c14n",
                    "malformed-refusal",
                ],
                "office": [
                    "docx-save-load",
                    "pptx-save-load",
                    "xlsx-save-load",
                    "aelira-pptx-reading-order-validate",
                ],
            },
            sort_keys=True,
        )
    )
